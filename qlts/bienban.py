"""Biên bản nghiệm thu / bàn giao (mẫu MH/QT-02/M04): xuất Word, PDF, Excel.

Mỗi biên bản (``BienBan``) là một trang (PDF/Word) hoặc một sheet (Excel).
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pandas as pd

from .config import app_setting

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "bien_ban_ban_giao.docx"
LOGO = ROOT / "assets" / "bienban_logo.jpeg"
FONTS = ROOT / "assets" / "fonts"
HEADERS = ["STT", "TÊN HÀNG HÓA/DỊCH VỤ", "MÃ HIỆU/ QUY CÁCH", "ĐVT", "SỐ LƯỢNG", "GHI CHÚ"]


@dataclass
class BienBan:
    giao_ten: str
    giao_chuc_vu: str
    nhan_ten: str
    nhan_chuc_vu: str
    items: list[dict]  # ten, quy_cach, dvt, sl, ghi_chu
    ngay: date = field(default_factory=date.today)
    dia_diem: str = ""
    don_vi: str = ""  # Đơn vị sử dụng tài sản (biên bản bàn giao theo phòng)
    sheet: str = "Bien ban"

    @property
    def place(self) -> str:
        return self.dia_diem or default_place()

    def date_line(self) -> str:
        return (f"Hôm nay, ngày {self.ngay.day:02d} tháng {self.ngay.month:02d} năm {self.ngay.year}, "
                f"tại {self.place} chúng tôi gồm:")

    def receiver_heading(self) -> str:
        return f"Đại diện bên nhận – Đơn vị sử dụng tài sản: {self.don_vi}" if self.don_vi else "Đại diện bên nhận:"


def default_place() -> str:
    return str(app_setting("bienban_dia_diem", "Văn phòng Trường TH, THCS và THPT Tân Phú"))


def ma_so() -> str:
    return str(app_setting("bienban_ma_so", "MH/QT-02/M04"))


def hieu_luc() -> str:
    return str(app_setting("bienban_hieu_luc", "14/7/2025"))


def _fmt_qty(value: float) -> str:
    return f"{value:g}".replace(".", ",")


def group_items(df: pd.DataFrame, note_code: bool = True) -> list[dict]:
    """Gom thiết bị cùng Mã tài sản (và cùng quy cách): Quy cách = Chi tiết, số lượng cộng lại."""
    if df.empty:
        return []
    work = df.copy()
    # Quy cách = Chi tiết; nếu Chi tiết trống hoặc trùng Tên thiết bị thì dùng Đặc điểm
    same = work["ChiTiet"].str.strip().str.lower() == work["TenThietBi"].str.strip().str.lower()
    work["QuyCach"] = work["ChiTiet"].where((work["ChiTiet"].str.strip() != "") & ~same, work["DacDiem"])
    work["SLx"] = pd.to_numeric(work["SL"], errors="coerce").fillna(0)
    work.loc[work["SLx"] <= 0, "SLx"] = 1
    work["Ten"] = work["TenThietBi"].where(work["TenThietBi"].str.strip() != "", work["ChiTiet"])
    rows = []
    for (ma, ten, qc, dvt), grp in work.groupby(["MaTaiSan", "Ten", "QuyCach", "DVT"], sort=True, dropna=False):
        rows.append({"ten": ten, "quy_cach": qc, "dvt": dvt, "sl": float(grp["SLx"].sum()),
                     "ghi_chu": f"Mã TS: {ma}" if note_code and ma else ""})
    return rows


# ---------------------------------------------------------------------------
# Word (điền vào đúng file mẫu)
# ---------------------------------------------------------------------------
def _set_text(paragraph, text: str) -> None:
    """Đổi nội dung đoạn văn, giữ định dạng của run đầu tiên."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _set_labeled(paragraph, prefix: str, text: str) -> None:
    """Đổi đoạn dạng '<nhãn đậm> <nội dung>' (vd. '2. Đại diện bên nhận:'): giữ run nhãn đậm."""
    runs = paragraph.runs
    bold = [r for r in runs if r.bold]
    if not bold:
        _set_text(paragraph, prefix + text)
        return
    for r in runs:
        if r is not bold[0]:
            r._r.getparent().remove(r._r)
    bold[0].text = prefix + text


def _fill_body(doc, elements: list, bb: BienBan) -> None:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    paras = [Paragraph(e, doc) for e in elements if e.tag.endswith("}p")]
    tables = [Table(e, doc) for e in elements if e.tag.endswith("}tbl")]
    people = [(bb.giao_ten, bb.giao_chuc_vu), (bb.nhan_ten, bb.nhan_chuc_vu)]
    for p in paras:
        text = p.text.strip()
        if text.startswith("Hôm nay"):
            _set_text(p, bb.date_line())
        elif text.startswith("Họ và tên") and people:
            name, title = people.pop(0)
            _set_text(p, f"Họ và tên: {name}\tChức vụ: {title}")
        elif "Đại diện bên nhận" in text:
            _set_labeled(p, "2.  ", bb.receiver_heading())
    # Bảng tiêu đề: mã số / hiệu lực
    if tables:
        cell = tables[0].rows[0].cells[-1]
        if len(cell.paragraphs) >= 2:
            _set_text(cell.paragraphs[0], f"Mã số: {ma_so()}")
            _set_text(cell.paragraphs[1], f"Hiệu lực: {hieu_luc()}")
    # Bảng nội dung bàn giao
    if len(tables) >= 2:
        tbl = tables[1]
        proto = copy.deepcopy(tbl.rows[1]._tr)
        for row in list(tbl.rows)[1:]:
            tbl._tbl.remove(row._tr)
        for i, it in enumerate(bb.items, start=1):
            tr = copy.deepcopy(proto)
            tbl._tbl.append(tr)
            cells = tbl.rows[-1].cells
            values = [str(i), it["ten"], it["quy_cach"], it["dvt"], _fmt_qty(it["sl"]), it.get("ghi_chu", "")]
            for c, v in zip(cells, values):
                _set_text(c.paragraphs[0], v)
                for extra in c.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
    # Bảng chữ ký
    if len(tables) >= 3:
        sig = tables[2].rows[-1].cells
        for cell, name in ((sig[0], bb.giao_ten), (sig[-1], bb.nhan_ten)):
            _set_text(cell.paragraphs[-1], name)


def to_docx(docs: list[BienBan]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document(str(TEMPLATE))
    body = doc.element.body
    original = [e for e in body.iterchildren() if not e.tag.endswith("}sectPr")]
    sect = body.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
    for e in original:
        body.remove(e)
    for n, bb in enumerate(docs):
        block = [copy.deepcopy(e) for e in original]
        if n:
            brk = doc.add_paragraph()._p  # sang trang mới cho biên bản tiếp theo
            body.remove(brk)
            from docx.text.paragraph import Paragraph
            Paragraph(brk, doc).add_run().add_break(WD_BREAK.PAGE)
            block.insert(0, brk)
        _fill_body(doc, block, bb)
        for e in block:
            if sect is not None:
                sect.addprevious(e)
            else:
                body.append(e)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF (vẽ lại theo bố cục mẫu; mỗi biên bản bắt đầu một trang mới)
# ---------------------------------------------------------------------------
def to_pdf(docs: list[BienBan]) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    pdf = FPDF(format="A4", unit="mm")
    pdf.set_margins(25, 12, 15)
    pdf.set_auto_page_break(True, 15)
    pdf.add_font("TNR", "", str(FONTS / "LiberationSerif-Regular.ttf"))
    pdf.add_font("TNR", "B", str(FONTS / "LiberationSerif-Bold.ttf"))
    pdf.add_font("TNR", "I", str(FONTS / "LiberationSerif-Italic.ttf"))
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def line(text: str, style: str = "", indent: float = 0, h: float = 6.5, align: str = "L"):
        pdf.set_font("TNR", style, 12)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, h, text, align=align, new_x="LMARGIN", new_y="NEXT")

    def two_cols(left: str, right: str, indent: float = 8):
        pdf.set_font("TNR", "", 12)
        pdf.set_x(pdf.l_margin + indent)
        pdf.cell(95 - indent, 6.5, left)
        pdf.multi_cell(width - 95, 6.5, right, new_x="LMARGIN", new_y="NEXT")

    def checkbox(label: str):
        pdf.set_font("TNR", "", 12)
        x = pdf.l_margin + 8
        pdf.rect(x, pdf.get_y() + 1.5, 3.5, 3.5)
        pdf.set_x(x + 6)
        pdf.cell(0, 6.5, label, new_x="LMARGIN", new_y="NEXT")

    for bb in docs:
        pdf.add_page()
        # Tiêu đề: logo | tên biên bản | mã số
        y0, h0 = pdf.get_y(), 20
        w1, w3 = 32, 42
        w2 = width - w1 - w3
        pdf.set_line_width(0.3)
        pdf.rect(pdf.l_margin, y0, width, h0)
        pdf.line(pdf.l_margin + w1, y0, pdf.l_margin + w1, y0 + h0)
        pdf.line(pdf.l_margin + w1 + w2, y0, pdf.l_margin + w1 + w2, y0 + h0)
        if LOGO.exists():
            pdf.image(str(LOGO), x=pdf.l_margin + 3, y=y0 + 2, h=h0 - 4)
        pdf.set_font("TNR", "B", 14)
        pdf.set_xy(pdf.l_margin + w1, y0)
        pdf.cell(w2, h0, "BIÊN BẢN NGHIỆM THU/BÀN GIAO", align="C")
        pdf.set_font("TNR", "", 10)
        pdf.set_xy(pdf.l_margin + w1 + w2 + 2, y0 + 5)
        pdf.cell(w3 - 2, 5, f"Mã số: {ma_so()}", new_x="LEFT", new_y="NEXT")
        pdf.cell(w3 - 2, 5, f"Hiệu lực: {hieu_luc()}")
        pdf.set_y(y0 + h0 + 6)

        line(bb.date_line(), align="J")
        line("1. Đại diện bên giao:", "B")
        two_cols(f"Họ và tên: {bb.giao_ten}", f"Chức vụ: {bb.giao_chuc_vu}")
        line("2. " + bb.receiver_heading(), "B")
        two_cols(f"Họ và tên: {bb.nhan_ten}", f"Chức vụ: {bb.nhan_chuc_vu}")
        line("3. Nội dung bàn giao:", "B")
        pdf.ln(1)

        pdf.set_font("TNR", "", 11)
        head = FontFace(emphasis="BOLD", size_pt=10)
        left = pdf.l_margin
        with pdf.table(width=width, align="LEFT", col_widths=(10, 42, 56, 13, 16, 27), text_align=("C", "L", "L", "C", "C", "L"),
                       headings_style=head, line_height=5.5, padding=1) as table:
            row = table.row()
            for h in HEADERS:
                row.cell(h, align="C")
            for i, it in enumerate(bb.items, start=1):
                row = table.row()
                for v in (str(i), it["ten"], it["quy_cach"], it["dvt"], _fmt_qty(it["sl"]), it.get("ghi_chu", "")):
                    row.cell(str(v))
        pdf.set_left_margin(left)
        pdf.ln(3)

        line("4. Ý kiến của đơn vị nhận và đơn vị chuyên môn (nếu có):", "B")
        line("- Đơn vị nhận: " + "." * 110, indent=8)
        line("- Đơn vị chuyên môn (nếu có): " + "." * 85, indent=8)
        line("5. Kết quả nghiệm thu/bàn giao:", "B")
        checkbox("Đồng ý nghiệm thu/bàn giao;")
        checkbox("Không đồng ý nghiệm thu/bàn giao.")
        if pdf.get_y() > pdf.h - 70:  # giữ đoạn cuối và chữ ký trên cùng một trang
            pdf.add_page()
        line("Biên bản này được lập thành 02 (hai) bản, mỗi bên giữ 01 (một) bản có giá trị như nhau, "
             "đại diện các bên tham gia ký tên:", align="J")
        pdf.ln(2)
        col = width / 3
        pdf.set_font("TNR", "B", 12)
        y = pdf.get_y()
        for i, title in enumerate(("Bên giao", "Đơn vị chuyên môn\n(nếu có)", "Bên nhận")):
            pdf.set_xy(pdf.l_margin + i * col, y)
            pdf.multi_cell(col, 6, title, align="C")
        pdf.set_y(y + 38)
        for i, name in ((0, bb.giao_ten), (2, bb.nhan_ten)):
            pdf.set_xy(pdf.l_margin + i * col, y + 38)
            pdf.multi_cell(col, 6, name, align="C")
    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Excel (mỗi biên bản một sheet, đã căn trang A4 để in)
# ---------------------------------------------------------------------------
def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name).strip("'") or "Sheet"
    base = base[:31]
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(out.lower())
    return out


def to_xlsx(docs: list[BienBan]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    f = lambda **kw: Font(name="Times New Roman", size=kw.pop("size", 12), **kw)  # noqa: E731
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for bb in docs:
        ws = wb.create_sheet(_sheet_name(bb.sheet, used))
        for col, w in zip("ABCDEF", (6, 26, 36, 8, 10, 18)):
            ws.column_dimensions[col].width = w

        def put(row: int, text: str, bold: bool = False, cols: str = "A:F", align=wrap, height: float | None = None):
            a, b = cols.split(":")
            ws.merge_cells(f"{a}{row}:{b}{row}")
            cell = ws[f"{a}{row}"]
            cell.value, cell.font, cell.alignment = text, f(bold=bold), align
            if height:
                ws.row_dimensions[row].height = height

        # Tiêu đề
        ws.merge_cells("A1:B3")
        ws.merge_cells("C1:D3")
        ws.merge_cells("E1:F3")
        ws["C1"].value, ws["C1"].font, ws["C1"].alignment = "BIÊN BẢN NGHIỆM THU/BÀN GIAO", f(bold=True, size=14), center
        ws["E1"].value = f"Mã số: {ma_so()}\nHiệu lực: {hieu_luc()}"
        ws["E1"].font, ws["E1"].alignment = f(size=10), wrap
        for r in range(1, 4):
            ws.row_dimensions[r].height = 18
            for c in "ABCDEF":
                ws[f"{c}{r}"].border = box
        if LOGO.exists():
            img = XLImage(str(LOGO))
            img.height, img.width = 60, 60 * 159 / 99
            ws.add_image(img, "A1")

        r = 5
        put(r, bb.date_line(), height=32); r += 1
        put(r, "1. Đại diện bên giao:", bold=True); r += 1
        put(r, f"Họ và tên: {bb.giao_ten}", cols="A:C"); put(r, f"Chức vụ: {bb.giao_chuc_vu}", cols="D:F"); r += 1
        put(r, "2. " + bb.receiver_heading(), bold=True); r += 1
        put(r, f"Họ và tên: {bb.nhan_ten}", cols="A:C"); put(r, f"Chức vụ: {bb.nhan_chuc_vu}", cols="D:F"); r += 1
        put(r, "3. Nội dung bàn giao:", bold=True); r += 1

        for c, h in zip("ABCDEF", HEADERS):
            cell = ws[f"{c}{r}"]
            cell.value, cell.font, cell.alignment, cell.border = h, f(bold=True, size=10), center, box
        ws.row_dimensions[r].height = 30
        r += 1
        for i, it in enumerate(bb.items, start=1):
            values = (i, it["ten"], it["quy_cach"], it["dvt"], it["sl"], it.get("ghi_chu", ""))
            for c, v in zip("ABCDEF", values):
                cell = ws[f"{c}{r}"]
                cell.value, cell.font, cell.border = v, f(size=11), box
                cell.alignment = center if c in "ADE" else wrap
            r += 1
        r += 1
        put(r, "4. Ý kiến của đơn vị nhận và đơn vị chuyên môn (nếu có):", bold=True); r += 1
        put(r, "- Đơn vị nhận: " + "." * 100); r += 1
        put(r, "- Đơn vị chuyên môn (nếu có): " + "." * 80); r += 1
        put(r, "5. Kết quả nghiệm thu/bàn giao:", bold=True); r += 1
        put(r, "     □ Đồng ý nghiệm thu/bàn giao;"); r += 1
        put(r, "     □ Không đồng ý nghiệm thu/bàn giao."); r += 1
        put(r, "Biên bản này được lập thành 02 (hai) bản, mỗi bên giữ 01 (một) bản có giá trị như nhau, "
               "đại diện các bên tham gia ký tên:", height=32); r += 2
        put(r, "Bên giao", True, "A:B", center); put(r, "Đơn vị chuyên môn\n(nếu có)", True, "C:D", center, 32)
        put(r, "Bên nhận", True, "E:F", center)
        r += 6
        put(r, bb.giao_ten, True, "A:B", center); put(r, bb.nhan_ten, True, "E:F", center)

        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left, ws.page_margins.right = 0.6, 0.4
        ws.print_options.horizontalCentered = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
