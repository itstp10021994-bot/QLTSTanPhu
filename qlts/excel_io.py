"""Xuất / nhập Excel cho các SharePoint List: biểu mẫu trống, xuất dữ liệu, đọc file tải lên."""

from __future__ import annotations

import io
import re
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import schema

TYPE_NAMES = {
    "text": "Một dòng văn bản (Single line of text)",
    "note": "Nhiều dòng văn bản (Multiple lines of text)",
    "number": "Số (Number)",
    "date": "Ngày (Date and time – chỉ ngày)",
    # Tình trạng / Vai trò để dạng văn bản cho linh hoạt (dữ liệu cũ có nhiều giá trị khác nhau)
    "choice": "Một dòng văn bản (Single line of text)",
}

SAMPLES = {
    schema.THIET_BI: {
        "STT": 1, "MaTaiSan": "111028", "MaChiTiet": "111028-00001", "ChiTiet": "Máy chiếu",
        "DacDiem": "Epson, 1024x768, EB-530", "TenPhongBan": "VFQF610018L", "DVT": "Cái", "SL": 1,
        "NoiSuDung": "Trệt_PH001", "NguoiSuDung": "giaovien@truong.edu.vn", "NgayMua": date(2023, 8, 22),
        "QuanLyThietBi": "Lê Thiên An", "TinhTrang": "Bình thường", "QuanLyPhong": "giaovien@truong.edu.vn",
        "PhanQuyenTB": "admin", "MaSAP": "T000816001", "ThoiHanBaoHanh": "24 tháng", "GhiChu": "",
        "TenThietBi": "Máy chiếu", "GiaTri": 12000000, "NgayHoaDon": date(2023, 8, 22),
        "NhomThietBi": "Nhóm Công Nghệ Thông Tin", "Mail": "",
    },
    schema.PHONG: {
        "Title": "Trệt_PH001", "TenPhong": "Phòng học 001",
        "NguoiQuanLy": "giaovien@truong.edu.vn", "TenNguoiQuanLy": "Nguyễn Văn A",
    },
    schema.DIEU_CHUYEN: {
        "Title": "111028-00001", "TenThietBi": "Máy chiếu", "TuPhong": "Kho_TB", "DenPhong": "Trệt_PH001",
        "SoLuong": 1, "NgayDieuChuyen": date(2026, 9, 1), "NguoiThucHien": "qlts@truong.edu.vn",
        "LyDo": "Cấp cho phòng học",
    },
    schema.KIEM_KE: {
        "MaTaiSan": "111042", "STT": 1, "TenTaiSan": "Tivi", "DacDiem": "Sharp, 58 inch, UE630X", "DVT": "Cái",
        "SoLuong": 1, "SoLuongKiemKe": 1, "GhiChu": "", "QuanLyThietBi": "Lê Thiên An", "NoiSuDung": "Bảo vệ",
        "QuanLyPhong": "Nguyễn Trọng Hòa", "TrangThai": "Bình thường", "NgayKiemKe": date(2025, 11, 5),
        "TrangThaiKiemKe": "Đã kiểm kê", "NgayXacNhan": None, "TrangThaiXacNhan": "",
        "DotKiemKe": "Đợt 2 năm 2024 - 2025", "NguoiKiemKe": "kiemke@truong.edu.vn", "NguoiXacNhan": "",
    },
    schema.LOAI_TB: {"TenThietBi": "Màn hình LED", "MaPhanLoai": "111305", "NhomThietBi": "Nhóm Công Nghệ Thông Tin"},
    schema.PHAN_QUYEN: {
        "Title": "email_cua_ban@truong.edu.vn", "HoTen": "Lê Thiên An", "VaiTro": schema.ROLE_ADMIN,
        "ChucDanh": "Chuyên viên Quản lý hệ thống",
    },
}

NOTES = {
    "MaChiTiet": "Mã tài sản + '-' + 5 chữ số, app tự sinh khi nhập mới",
    "NoiSuDung": "Mã phòng (trùng cột Mã phòng của list Phong); 'Thanh lý' = đã đưa đi thanh lý",
    "QuanLyPhong": "Email người quản lý phòng – người này thấy phòng ở mục Người dùng",
    "TinhTrang": "Gợi ý: " + ", ".join(schema.TINH_TRANG),
    "VaiTro": "Một trong: " + ", ".join(schema.ROLES) + " (mỗi vai trò một dòng)",
    "SoLuong": "Số lượng sổ sách = tổng thiết bị cùng Mã tài sản + Đặc điểm + Nơi sử dụng",
    "TrangThaiKiemKe": "Đang kiểm kê / Đã kiểm kê (app ghi khi bấm Xác nhận kiểm kê)",
    "TrangThaiXacNhan": "Đã xác nhận / Không đồng ý (quản lý phòng xác nhận)",
    "NguoiKiemKe": "Không bắt buộc – có cột thì app ghi email người kiểm kê",
    "NguoiXacNhan": "Không bắt buộc – có cột thì app ghi email người xác nhận",
    "MaPhanLoai": "Mã tài sản của loại thiết bị – chọn tên khi nhập mới, app tự điền mã này",
    "Title": "",
}

LIST_TITLES = {
    schema.THIET_BI: "Thiết bị chi tiết",
    schema.PHONG: "Danh mục phòng",
    schema.DIEU_CHUYEN: "Lịch sử điều chuyển",
    schema.KIEM_KE: "Kiểm kê",
    schema.PHAN_QUYEN: "Phân quyền",
    schema.LOAI_TB: "Danh mục loại thiết bị",
}

FILES = {
    schema.THIET_BI: ("1_ThietBi.xlsx", "Data_Thietbichitiet (hoặc tên bạn chọn)"),
    schema.PHONG: ("2_Phong.xlsx", "Phong"),
    schema.DIEU_CHUYEN: ("3_DieuChuyen.xlsx", "DieuChuyen"),
    schema.KIEM_KE: ("4_KiemKe.xlsx", "Data_Thietbi"),
    schema.PHAN_QUYEN: ("5_PhanQuyen.xlsx", "PhanQuyen"),
    schema.LOAI_TB: ("6_LoaiThietBi.xlsx", "Data_Loaithietbi"),
}


def template_workbook(list_key: str, list_name: str, rows: pd.DataFrame | None = None) -> bytes:
    """File Excel có bảng tbl_<list> (biểu mẫu + sheet hướng dẫn).

    ``rows`` = None -> một dòng mẫu; ngược lại ghi các dòng dữ liệu (khóa cột của app)."""
    cols = schema.LISTS[list_key]["columns"]
    wb = Workbook()
    ws = wb.active
    ws.title = "DuLieu"
    ws.append([spec["label"] for spec in cols.values()])
    records = [SAMPLES[list_key]] if rows is None else rows.to_dict("records")
    for rec in records or [{}]:
        ws.append([_cell_value(spec["type"], rec.get(key)) for key, spec in cols.items()])
    for i, (key, spec) in enumerate(cols.items(), start=1):
        for row_cells in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            if spec["type"] == "date":
                row_cells[0].number_format = "DD/MM/YYYY"
            elif spec["type"] == "number":
                row_cells[0].number_format = "#,##0"
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(spec["label"]) + 4)
    table = Table(displayName=f"tbl_{list_key}", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    guide = wb.create_sheet("HuongDan")
    guide.append([f"List SharePoint: {list_name}"])
    guide["A1"].font = Font(bold=True, size=13)
    guide.append(["Tạo list: Microsoft Lists → + Danh sách mới → Từ Excel → chọn file này → chọn bảng "
                  f"tbl_{list_key} → kiểm tra kiểu cột theo bảng dưới → Tạo. Sau đó XÓA dòng mẫu."])
    guide.append(["Không đổi tên cột: app dò cột theo đúng tên này."])
    guide.append([])
    guide.append(["Tên cột", "Kiểu cột khi tạo list", "Ghi chú"])
    for c in guide[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E8C")
    for key, spec in cols.items():
        guide.append([spec["label"], TYPE_NAMES[spec["type"]], NOTES.get(key, "")])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 44
    guide.column_dimensions["C"].width = 70
    for row in guide.iter_rows(min_row=2, max_row=3):
        row[0].alignment = Alignment(wrap_text=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()




def _cell_value(kind: str, value):
    if value is None or (isinstance(value, float) and pd.isna(value)) or value == "":
        return None
    if kind == "date":
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return str(value)
    if kind == "number":
        return float(value) if not float(value).is_integer() else int(value)
    return value


# ---------------------------------------------------------------------------
# Đọc file tải lên
# ---------------------------------------------------------------------------
def parse_number(value):
    """'12000000', '12000000.0', '12.000.000', '12,000,000' -> 12000000."""
    text = str(value).strip().replace(" ", "")
    num = pd.to_numeric(text, errors="coerce")
    if pd.isna(num):
        num = pd.to_numeric(text.replace(".", "").replace(",", ""), errors="coerce")
    return None if pd.isna(num) else float(num)


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def read_upload(file, list_key: str) -> tuple[pd.DataFrame, list[str]]:
    """Đọc .xlsx/.csv; cột khớp theo nhãn tiếng Việt hoặc khóa. Trả về (dữ liệu theo khóa app, cột bỏ qua)."""
    raw = pd.read_csv(file, dtype=str) if file.name.lower().endswith(".csv") else pd.read_excel(file, dtype=str)
    lookup = {_norm(label): key for key, label in schema.labels_of(list_key).items()}
    lookup.update({_norm(k): k for k in schema.columns_of(list_key)})
    rename, unknown = {}, []
    for col in raw.columns:
        key = lookup.get(_norm(col))
        if key and key not in rename.values():
            rename[col] = key
        else:
            unknown.append(str(col))
    data = raw[list(rename)].rename(columns=rename).dropna(how="all")
    records = []
    for rec in data.to_dict("records"):
        clean = {}
        for k, v in rec.items():
            if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
                continue
            v = str(v).strip()
            kind = schema.column_type(list_key, k)
            if kind == "number":
                v = parse_number(v)
            elif kind == "date":
                d = pd.to_datetime(v, dayfirst=not re.match(r"^\d{4}-", v), errors="coerce")
                v = d.date() if not pd.isna(d) else v
            clean[k] = v
        records.append(clean)
    return pd.DataFrame(records, columns=list(rename.values())), unknown


# Khóa để nhận ra dòng đã có (cập nhật thay vì thêm mới)
MATCH_KEYS = {
    schema.THIET_BI: ["MaChiTiet"],
    schema.PHONG: ["Title"],
    schema.PHAN_QUYEN: ["Title", "VaiTro"],
    schema.KIEM_KE: ["DotKiemKe", "NoiSuDung", "MaTaiSan", "DacDiem"],
    schema.DIEU_CHUYEN: [],  # lịch sử: luôn thêm mới
    schema.LOAI_TB: ["MaPhanLoai"],
}


def row_key(list_key: str, row) -> tuple | None:
    keys = MATCH_KEYS[list_key]
    if not keys:
        return None
    values = tuple(_norm(row.get(k, "") if isinstance(row, dict) else row[k]) for k in keys)
    return values if all(values) else None


REQUIRED = {
    schema.PHONG: ["Title"],
    schema.PHAN_QUYEN: ["Title", "VaiTro"],
    schema.KIEM_KE: ["DotKiemKe", "NoiSuDung", "MaTaiSan"],
    schema.DIEU_CHUYEN: ["Title"],
    schema.THIET_BI: [],  # kiểm tra riêng (cần Mã tài sản hoặc Mã chi tiết)
    schema.LOAI_TB: ["MaPhanLoai", "TenThietBi"],
}


def plan_import(list_key: str, data: pd.DataFrame, existing: pd.DataFrame, update_existing: bool,
                user_name: str = "", managers: dict | None = None) -> dict:
    """Lập danh sách thao tác ghi cho dữ liệu nhập.

    Trả về {"ops": [...], "create": n, "update": n, "skip": n, "problems": [...]}; số dòng Excel tính từ 2.
    """
    from . import thietbi

    labels = schema.labels_of(list_key)
    index = {}
    for rec in existing.to_dict("records"):
        k = row_key(list_key, rec)
        if k:
            index.setdefault(k, rec["id"])
    ops, problems = [], []
    n_create = n_update = n_skip = 0
    seen: dict = {}
    used_codes: dict = {}
    stt = [int(existing["STT"].max()) if "STT" in existing and not existing.empty else 0]
    codes = list(existing["MaChiTiet"]) if list_key == schema.THIET_BI else []

    for n, rec in enumerate(data.to_dict("records"), start=2):
        rec = {k: v for k, v in rec.items() if not (v is None or (isinstance(v, float) and pd.isna(v)))}
        for email_col in {schema.PHAN_QUYEN: "Title", schema.PHONG: "NguoiQuanLy"}.get(list_key, "").split():
            if email_col in rec:
                rec[email_col] = str(rec[email_col]).strip().lower()
        missing = [labels[k] for k in REQUIRED[list_key] if not rec.get(k)]
        if missing:
            problems.append(f"Dòng {n}: thiếu {', '.join(missing)}")
            continue
        k = row_key(list_key, rec)
        if k and k in seen:
            problems.append(f"Dòng {n}: trùng với dòng {seen[k]} trong file")
            continue
        if k:
            seen[k] = n
        if k and k in index:
            if update_existing:
                ops.append(("update", index[k], rec))
                n_update += 1
            else:
                n_skip += 1
            continue
        if list_key == schema.THIET_BI:
            rec = thietbi.prepare_new(rec, codes, used_codes, stt, managers or {}, user_name)
            if isinstance(rec, str):
                problems.append(f"Dòng {n}: {rec}")
                continue
        elif list_key == schema.KIEM_KE and not rec.get("STT"):
            stt[0] += 1
            rec["STT"] = stt[0]
        ops.append(("create", rec))
        n_create += 1
    return {"ops": ops, "create": n_create, "update": n_update, "skip": n_skip, "problems": problems}
