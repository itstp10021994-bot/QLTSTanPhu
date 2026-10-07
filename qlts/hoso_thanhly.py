"""Hồ sơ sau khi xác nhận thanh lý một đợt – 3 mẫu của quy trình HC/QT–02:

- M02: Đề nghị mang tài sản ra ngoài
- M06: Biên bản bàn giao – thanh lý tài sản (bên mua đã thanh toán đủ trước khi nhận)
- M07: Báo cáo kết quả thanh lý tài sản

Mỗi mẫu xuất Word (soạn lại theo bố cục file mẫu: khung LOGO | tên biểu mẫu | mã số) và PDF.
Font Times New Roman (PDF dùng Liberation Serif – cùng kích thước chữ).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date

import pandas as pd

from .bienban import LOGO
from .config import app_setting
from .thanhly import _pdf, money, qty

FORMS = {
    "M02": "ĐỀ NGHỊ MANG TÀI SẢN\nRA NGOÀI",
    "M06": "BIÊN BẢN BÀN GIAO -\nTHANH LÝ TÀI SẢN",
    "M07": "BÁO CÁO KẾT QUẢ THANH LÝ TÀI SẢN",
}
FONT = "Times New Roman"


def form_code(form: str) -> str:
    return "\n".join([
        f"Mã số: {app_setting(f'thanhly_{form.lower()}_ma_so', f'HC/QT–02/{form}')}",
        f"Lần ban hành: {app_setting('thanhly_lan_ban_hanh', '08')}",
        f"Hiệu lực: {app_setting('thanhly_hieu_luc', '25/11/2025')}",
    ])


def school() -> str:
    return str(app_setting("school_full_name", "Trường TH, THCS và THPT Tân Phú"))


def school_address() -> str:
    return str(app_setting("school_address", ""))


def dmy(d: date | None) -> str:
    return f"{d:%d/%m/%Y}" if d else "....../....../........."


# ---------------------------------------------------------------------------
# Đọc số tiền bằng chữ
# ---------------------------------------------------------------------------
_DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]


def _three(n: int, full: bool) -> list[str]:
    h, t, u = n // 100, n // 10 % 10, n % 10
    words: list[str] = []
    if h or full:
        words += [_DIGITS[h], "trăm"]
    if t == 0:
        if u and (h or full):
            words.append("lẻ")
    elif t == 1:
        words.append("mười")
    else:
        words += [_DIGITS[t], "mươi"]
    if u:
        if u == 1 and t > 1:
            words.append("mốt")
        elif u == 5 and t > 0:
            words.append("lăm")
        elif u == 4 and t > 1:
            words.append("tư")
        else:
            words.append(_DIGITS[u])
    return words


def so_thanh_chu(value) -> str:
    """1250000 -> 'Một triệu hai trăm năm mươi nghìn'."""
    try:
        n = int(round(float(value or 0)))
    except (TypeError, ValueError):
        return ""
    if n == 0:
        return "Không"
    units = ["", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ"]
    groups = []
    while n:
        groups.append(n % 1000)
        n //= 1000
    words: list[str] = []
    for i in range(len(groups) - 1, -1, -1):
        if groups[i] == 0:  # nhóm 0 (vd 1.000.500 -> "một triệu năm trăm"): bỏ qua
            continue
        words += _three(groups[i], full=i < len(groups) - 1)
        if units[i]:
            words.append(units[i])
    text = " ".join(words)
    return text[:1].upper() + text[1:]


# ---------------------------------------------------------------------------
# Dữ liệu từng mẫu
# ---------------------------------------------------------------------------
def group_items(items: pd.DataFrame) -> pd.DataFrame:
    """Gộp các dòng cùng Tên + Đặc điểm + ĐVT (cộng số lượng) – dùng cho M02, M06."""
    if items.empty:
        return pd.DataFrame(columns=["ten", "dac_diem", "dvt", "sl", "ma"])
    df = items.assign(sl=pd.to_numeric(items["sl"], errors="coerce").fillna(0))
    g = (df.groupby(["ten", "dac_diem", "dvt"], sort=False, dropna=False)
         .agg(sl=("sl", "sum"), ma=("ma", lambda s: ", ".join(s.astype(str)))).reset_index())
    return g


@dataclass
class DeNghi:  # M02
    items: pd.DataFrame  # ten, dac_diem, dvt, sl, ghi_chu
    nguoi_de_nghi: str = ""
    don_vi: str = ""
    noi: str = ""  # [Công ty/Trường/Trung tâm]
    ly_do: str = ""
    ky_hieu_truong: str = ""
    ky_quan_ly: str = ""
    ngay: date = field(default_factory=date.today)


@dataclass
class BanGiaoTL:  # M06
    items: pd.DataFrame  # ten, dac_diem, dvt, sl, don_gia
    ngay: date = field(default_factory=date.today)
    ben_thanh_ly: str = ""
    dia_chi: str = ""
    dai_dien: list[tuple[str, str]] = field(default_factory=list)  # (họ tên, chức vụ)
    ben_mua: str = ""
    dia_chi_mua: str = ""
    cccd_mst: str = ""
    nguoi_dai_dien: str = ""
    cccd_dai_dien: str = ""
    so_tien: float | None = None  # để trống = tổng thành tiền
    ngay_thanh_toan: date | None = None

    @property
    def thanh_tien(self) -> pd.Series:
        sl = pd.to_numeric(self.items["sl"], errors="coerce").fillna(0)
        return sl * pd.to_numeric(self.items["don_gia"], errors="coerce").fillna(0)

    @property
    def tong(self) -> float:
        return float(self.thanh_tien.sum()) if not self.items.empty else 0.0

    @property
    def da_thanh_toan(self) -> float:
        return self.tong if self.so_tien is None else float(self.so_tien)


@dataclass
class KetQua:  # M07
    items: pd.DataFrame  # ngay_mua, ma, ten, dac_diem, dvt, sl, gia_mua, con_lai
    can_cu: str = ""  # vd "Tờ trình số: 491/2025/TTr-TP ngày 10 tháng 11 năm 2025 về việc thanh lý ..."
    nhan_su: list[tuple[str, str]] = field(default_factory=list)
    chi_phi: str = "không phát sinh"
    gia_tri_thu_hoi: float = 0.0
    ngay_ghi_giam: date | None = None
    nguoi_lap: str = ""
    ke_toan: str = ""
    quan_ly: str = ""
    tham_quyen: str = ""

    def thu_hoi_text(self) -> str:
        return (f"Giá trị thu hồi: {money(self.gia_tri_thu_hoi)} đồng (bằng chữ: "
                f"{so_thanh_chu(self.gia_tri_thu_hoi).lower() if self.gia_tri_thu_hoi else '......'} đồng), "
                "đã bao gồm VAT.")

    def ghi_giam_text(self) -> str:
        d = self.ngay_ghi_giam
        return (f"Đã ghi giảm tài sản: ngày {d.day} tháng {d.month} năm {d.year}." if d
                else "Đã ghi giảm tài sản: ngày……tháng..…năm……..")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def _new_doc(landscape: bool = False):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(3)
    return doc


def _para(doc_or_cell, text: str = "", bold: bool = False, italic: bool = False, size: float = 12,
          align: str = "left", label: str = ""):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = doc_or_cell.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    if label:
        r = p.add_run(label)
        r.bold, r.font.size = True, Pt(size)
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    return p


def _cell_text(cell, text: str, bold: bool = False, size: float = 11, align: str = "center", italic=False):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)


def _header(doc, form: str) -> None:
    from docx.shared import Cm

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    widths = [Cm(4.2), Cm(8.6), Cm(4.6)]
    for cell, w in zip(t.rows[0].cells, widths):
        cell.width = w
    if LOGO.exists():
        p = t.rows[0].cells[0].paragraphs[0]
        p.alignment = 1
        p.add_run().add_picture(str(LOGO), height=Cm(1.6))
    _cell_text(t.rows[0].cells[1], FORMS[form], bold=True, size=14)
    _cell_text(t.rows[0].cells[2], form_code(form), size=10)
    doc.add_paragraph()


def _grid(doc, headers: list[str], rows: list[list[str]], widths_cm: list[float], aligns: str,
          total: list[str] | None = None, size: float = 11):
    from docx.shared import Cm

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        _cell_text(t.rows[0].cells[i], h, bold=True, size=size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            _cell_text(cells[i], v, size=size, align={"C": "center", "L": "left", "R": "right"}[aligns[i]])
    if total:
        cells = t.add_row().cells
        for i, v in enumerate(total):
            _cell_text(cells[i], v, bold=True, size=size,
                       align={"C": "center", "L": "left", "R": "right"}[aligns[i]])
    for row in t.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
    return t


def _signatures(doc, blocks: list[tuple[str, str]], gap_lines: int = 4, size: float = 11.5) -> None:
    t = doc.add_table(rows=2, cols=len(blocks))
    for i, (title, name) in enumerate(blocks):
        _cell_text(t.rows[0].cells[i], title, bold=True, size=size)
        _cell_text(t.rows[1].cells[i], "\n" * gap_lines + name, size=size)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def m02_docx(d: DeNghi) -> bytes:
    doc = _new_doc()
    _header(doc, "M02")
    _para(doc, d.nguoi_de_nghi or "." * 50, label="Người đề nghị: ")
    _para(doc, d.don_vi or "." * 50, label="Đơn vị: ")
    _para(doc, f"Đề nghị được mang ra ngoài {d.noi or school()} các tài sản sau:")
    rows = [[str(i), r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]), r.get("ghi_chu", "")]
            for i, r in enumerate(d.items.to_dict("records"), start=1)]
    _grid(doc, ["Stt", "Tên tài sản", "Đặc điểm", "ĐVT", "Số lượng", "Ghi chú"], rows,
          [1.1, 4.3, 5.4, 1.5, 1.8, 3.3], "CLLCCL")
    _para(doc, d.ly_do or "." * 120, label="Lý do: ")
    doc.add_paragraph()
    _signatures(doc, [("Người đề nghị", d.nguoi_de_nghi), ("Xét duyệt của\nCTQ/ Hiệu trưởng", d.ky_hieu_truong),
                      ("Xác nhận của\nĐơn vị Quản lý tài sản", d.ky_quan_ly)])
    doc.add_paragraph()
    _para(doc, "Xác nhận của Đơn vị phụ trách bảo vệ (nếu có, không áp dụng trong trường hợp máy tính xách tay):",
          bold=True)
    _para(doc, "Đem ra ngày …./…../20…giờ…..., Chữ ký: ____________ Họ tên: ................................")
    _para(doc, "Đem vào ngày …./…../20…giờ…., Chữ ký: ____________ Họ tên: ................................")
    return _save(doc)


def m06_docx(b: BanGiaoTL) -> bytes:
    doc = _new_doc()
    _header(doc, "M06")
    _para(doc, f"Hôm nay, ngày {b.ngay.day} tháng {b.ngay.month} năm {b.ngay.year}, tại {b.ben_thanh_ly or school()}"
               f", địa chỉ: {b.dia_chi or '.' * 60}")
    _para(doc, "Chúng tôi gồm có:")
    _para(doc, (b.ben_thanh_ly or school()).upper(), label="BÊN THANH LÝ: ")
    _para(doc, "Đại diện bởi:")
    reps = b.dai_dien or [("", ""), ("", "")]
    for name, title in reps:
        _para(doc, f"Ông/ Bà: {name or '.' * 50}" + (f" – Chức vụ: {title}" if title else ""))
    _para(doc, b.ben_mua or "[Ông/Bà/Công ty]", label="BÊN MUA: ")
    _para(doc, f"Địa chỉ: {b.dia_chi_mua or '.' * 70}")
    _para(doc, f"CCCD/CMND/Mã số thuế: {b.cccd_mst or '.' * 50}")
    _para(doc, f"Người đại diện: {b.nguoi_dai_dien or '.' * 50}")
    _para(doc, f"CCCD/CMND của người đại diện: {b.cccd_dai_dien or '.' * 40}")
    _para(doc, f"Bên Mua cam kết đã hoàn tất nghĩa vụ thanh toán số tiền là {money(b.da_thanh_toan)} VNĐ vào ngày "
               f"{dmy(b.ngay_thanh_toan)} cho Bên Thanh Lý để nhận tài sản thanh lý, nay Bên Thanh Lý giao và Bên "
               "Mua chấp nhận tài sản thanh lý theo danh mục cụ thể dưới đây:", align="justify")
    tt = b.thanh_tien
    rows = [[str(i), r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]), money(r["don_gia"]), money(tt.iloc[i - 1])]
            for i, r in enumerate(b.items.to_dict("records"), start=1)]
    _grid(doc, ["Stt", "Tên tài sản", "Đặc điểm", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"], rows,
          [1.1, 3.8, 4.6, 1.4, 1.7, 2.4, 2.6], "CLLCCRR",
          total=["", "Tổng cộng", "", "", qty(pd.to_numeric(b.items["sl"], errors="coerce").sum()), "",
                 money(b.tong)])
    _para(doc, f"Tổng cộng (bằng số): {money(b.tong)} VNĐ.")
    _para(doc, f"Bằng chữ: {so_thanh_chu(b.tong) if b.tong else '......'} đồng Việt Nam.")
    _para(doc, "Biên bản này được lập thành 02 (hai) bản vào ngày "
               f"{dmy(b.ngay)}, mỗi bên giữ 01 (một) bản và có giá trị như nhau./.", align="justify")
    doc.add_paragraph()
    _signatures(doc, [("ĐẠI DIỆN BÊN MUA\n(ký và ghi rõ họ tên)", b.nguoi_dai_dien or ""),
                      ("ĐẠI DIỆN BÊN THANH LÝ\n(ký và ghi rõ họ tên)", reps[0][0] if reps else "")])
    return _save(doc)


def m07_docx(k: KetQua) -> bytes:
    doc = _new_doc(landscape=True)
    _header(doc, "M07")
    _para(doc, f"Thực hiện theo {k.can_cu or '……. số:....../20…/……. ngày.....tháng.....năm 20… về việc thanh lý ….'}",
          align="justify")
    _para(doc, "Các nhân sự phụ trách việc thanh lý tài sản gồm:")
    for name, title in (k.nhan_su or [("", "")] * 3):
        _para(doc, f"Ông/Bà {name or '.' * 40} - Chức vụ {title or '.' * 30}")
    _para(doc, "Tiến hành thanh lý tài sản", bold=True)
    rows = [[str(i), str(r.get("ngay_mua", "")), r["ma"], r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]),
             money(r["gia_mua"]), money(r["con_lai"])] for i, r in enumerate(k.items.to_dict("records"), start=1)]
    sl = pd.to_numeric(k.items["sl"], errors="coerce").fillna(0).sum()
    _grid(doc, ["STT", "Ngày mua", "Mã tài sản", "Tên tài sản", "Đặc điểm", "ĐVT", "SL", "Nguyên giá",
                "Giá còn lại"], rows, [1.2, 2.0, 3.6, 4.0, 5.6, 1.4, 1.3, 2.9, 2.9], "CCCLLCCRR",
          total=["", "", "", "Tổng cộng", "", "", qty(sl),
                 money(pd.to_numeric(k.items["gia_mua"], errors="coerce").sum()),
                 money(pd.to_numeric(k.items["con_lai"], errors="coerce").sum())], size=10.5)
    _para(doc, "Kết quả thanh lý tài sản", bold=True)
    _para(doc, f"Chi phí thanh lý tài sản: {k.chi_phi or 'không phát sinh'}.")
    _para(doc, k.thu_hoi_text())
    _para(doc, k.ghi_giam_text())
    doc.add_paragraph()
    _signatures(doc, [("NGƯỜI LẬP BÁO CÁO", k.nguoi_lap), ("ĐƠN VỊ PHỤ TRÁCH KẾ TOÁN", k.ke_toan),
                      ("ĐƠN VỊ QUẢN LÝ TÀI SẢN/\nHỘI ĐỒNG THANH LÝ TÀI SẢN", k.quan_ly),
                      ("CẤP THẨM QUYỀN", k.tham_quyen)])
    return _save(doc)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_header(pdf, form: str) -> None:
    width = pdf.w - pdf.l_margin - pdf.r_margin
    x0, y0, h0 = pdf.l_margin, pdf.get_y(), 20
    a, b = x0 + width * 0.24, x0 + width * 0.74
    pdf.rect(x0, y0, width, h0)
    pdf.line(a, y0, a, y0 + h0)
    pdf.line(b, y0, b, y0 + h0)
    if LOGO.exists():
        lw = (h0 - 5) * 159 / 99
        pdf.image(str(LOGO), x=x0 + (a - x0 - lw) / 2, y=y0 + 2.5, h=h0 - 5)
    pdf.set_font("TNR", "B", 14)
    lines = FORMS[form].split("\n")
    pdf.set_xy(a, y0 + (h0 - 7 * len(lines)) / 2)
    pdf.multi_cell(b - a, 7, FORMS[form], align="C")
    pdf.set_font("TNR", "", 10)
    pdf.set_xy(b, y0 + 3)
    pdf.multi_cell(x0 + width - b, 4.8, form_code(form), align="C")
    pdf.set_xy(x0, y0 + h0 + 5)


def _line(pdf, text: str, label: str = "", style: str = "", size: float = 12, h: float = 6.5) -> None:
    pdf.set_x(pdf.l_margin)
    if label:
        pdf.set_font("TNR", "B", size)
        pdf.write(h, label)
    pdf.set_font("TNR", style, size)
    pdf.write(h, text)
    pdf.ln(h + 0.5)


def _pdf_table(pdf, headers, rows, widths, aligns, total=None, size: float = 10.5) -> None:
    from fpdf.fonts import FontFace

    width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("TNR", "", size)
    with pdf.table(width=width, align="LEFT", col_widths=widths, text_align=tuple(aligns),
                   line_height=size * 0.5, headings_style=FontFace(emphasis="BOLD", size_pt=size), padding=1) as t:
        r = t.row()
        for h in headers:
            r.cell(h, align="C")
        for row in rows:
            r = t.row()
            for v in row:
                r.cell(str(v))
        if total:
            r = t.row(style=FontFace(emphasis="BOLD"))
            for v in total:
                r.cell(str(v))
    pdf.ln(2)


def _pdf_signatures(pdf, blocks: list[tuple[str, str]], size: float = 11.5) -> None:
    width = pdf.w - pdf.l_margin - pdf.r_margin
    if pdf.get_y() > pdf.h - pdf.b_margin - 31:  # giữ cả khối ký tên trên một trang
        pdf.add_page()
    y = pdf.get_y() + 3
    w = width / len(blocks)
    for i, (title, name) in enumerate(blocks):
        x = pdf.l_margin + i * w
        pdf.set_font("TNR", "B", size)
        pdf.set_xy(x, y)
        pdf.multi_cell(w, 5.5, title, align="C")
        pdf.set_font("TNR", "", size)
        pdf.set_xy(x, y + 22)
        pdf.cell(w, 6, name, align="C")
    pdf.set_y(y + 30)


def m02_pdf(d: DeNghi) -> bytes:
    pdf = _pdf("P")
    pdf.set_margins(18, 14, 16)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    _pdf_header(pdf, "M02")
    _line(pdf, d.nguoi_de_nghi or "." * 60, "Người đề nghị: ")
    _line(pdf, d.don_vi or "." * 60, "Đơn vị: ")
    _line(pdf, f"Đề nghị được mang ra ngoài {d.noi or school()} các tài sản sau:")
    rows = [[i, r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]), r.get("ghi_chu", "")]
            for i, r in enumerate(d.items.to_dict("records"), start=1)]
    _pdf_table(pdf, ["Stt", "Tên tài sản", "Đặc điểm", "ĐVT", "Số lượng", "Ghi chú"], rows,
               [10, 40, 52, 14, 17, 33], "CLLCCL")
    _line(pdf, d.ly_do or "." * 110, "Lý do: ")
    _pdf_signatures(pdf, [("Người đề nghị", d.nguoi_de_nghi), ("Xét duyệt của\nCTQ/ Hiệu trưởng", d.ky_hieu_truong),
                          ("Xác nhận của\nĐơn vị Quản lý tài sản", d.ky_quan_ly)])
    _line(pdf, "Xác nhận của Đơn vị phụ trách bảo vệ (nếu có, không áp dụng trong trường hợp máy tính xách tay):",
          style="B")
    pdf.ln(2)
    _line(pdf, "Đem ra ngày …./…../20…giờ…..., Chữ ký: ____________ Họ tên: .............................")
    pdf.ln(2)
    _line(pdf, "Đem vào ngày …./…../20…giờ…., Chữ ký: ____________ Họ tên: .............................")
    return bytes(pdf.output())


def m06_pdf(b: BanGiaoTL) -> bytes:
    pdf = _pdf("P")
    pdf.set_margins(18, 14, 16)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    _pdf_header(pdf, "M06")

    def para(text, label="", style=""):
        pdf.set_x(pdf.l_margin)
        if label:
            pdf.set_font("TNR", "B", 12)
            pdf.write(6.3, label)
        pdf.set_font("TNR", style, 12)
        pdf.multi_cell(0, 6.3, text, align="J" if len(text) > 90 else "L", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.5)

    para(f"Hôm nay, ngày {b.ngay.day} tháng {b.ngay.month} năm {b.ngay.year}, tại {b.ben_thanh_ly or school()}, "
         f"địa chỉ: {b.dia_chi or '.' * 60}")
    para("Chúng tôi gồm có:")
    para((b.ben_thanh_ly or school()).upper(), "BÊN THANH LÝ: ")
    para("Đại diện bởi:")
    reps = b.dai_dien or [("", ""), ("", "")]
    for name, title in reps:
        para(f"Ông/ Bà: {name or '.' * 50}" + (f" – Chức vụ: {title}" if title else ""))
    para(b.ben_mua or "[Ông/Bà/Công ty]", "BÊN MUA: ")
    para(f"Địa chỉ: {b.dia_chi_mua or '.' * 70}")
    para(f"CCCD/CMND/Mã số thuế: {b.cccd_mst or '.' * 50}")
    para(f"Người đại diện: {b.nguoi_dai_dien or '.' * 50}")
    para(f"CCCD/CMND của người đại diện: {b.cccd_dai_dien or '.' * 40}")
    para(f"Bên Mua cam kết đã hoàn tất nghĩa vụ thanh toán số tiền là {money(b.da_thanh_toan)} VNĐ vào ngày "
         f"{dmy(b.ngay_thanh_toan)} cho Bên Thanh Lý để nhận tài sản thanh lý, nay Bên Thanh Lý giao và Bên Mua "
         "chấp nhận tài sản thanh lý theo danh mục cụ thể dưới đây:")
    tt = b.thanh_tien
    rows = [[i, r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]), money(r["don_gia"]), money(tt.iloc[i - 1])]
            for i, r in enumerate(b.items.to_dict("records"), start=1)]
    _pdf_table(pdf, ["Stt", "Tên tài sản", "Đặc điểm", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"], rows,
               [10, 36, 44, 13, 16, 21, 26], "CLLCCRR",
               total=["", "Tổng cộng", "", "", qty(pd.to_numeric(b.items["sl"], errors="coerce").sum()), "",
                      money(b.tong)])
    para(f"Tổng cộng (bằng số): {money(b.tong)} VNĐ.")
    para(f"Bằng chữ: {so_thanh_chu(b.tong) if b.tong else '......'} đồng Việt Nam.")
    para(f"Biên bản này được lập thành 02 (hai) bản vào ngày {dmy(b.ngay)}, mỗi bên giữ 01 (một) bản và có giá trị "
         "như nhau./.")
    _pdf_signatures(pdf, [("ĐẠI DIỆN BÊN MUA\n(ký và ghi rõ họ tên)", b.nguoi_dai_dien),
                          ("ĐẠI DIỆN BÊN THANH LÝ\n(ký và ghi rõ họ tên)", reps[0][0] if reps else "")])
    return bytes(pdf.output())


def m07_pdf(k: KetQua) -> bytes:
    pdf = _pdf("L")
    pdf.set_margins(14, 12, 14)
    pdf.set_auto_page_break(True, 12)
    pdf.add_page()
    _pdf_header(pdf, "M07")
    pdf.set_y(pdf.get_y() - 2)
    pdf.set_font("TNR", "", 12)
    pdf.multi_cell(0, 6.3, "Thực hiện theo " + (k.can_cu or "……. số:....../20…/……. ngày.....tháng.....năm 20… về "
                                                "việc thanh lý ……………………."), new_x="LMARGIN", new_y="NEXT")
    _line(pdf, "Các nhân sự phụ trách việc thanh lý tài sản gồm:")
    for name, title in (k.nhan_su or [("", "")] * 3):
        _line(pdf, f"Ông/Bà {name or '.' * 40} - Chức vụ {title or '.' * 30}")
    _line(pdf, "Tiến hành thanh lý tài sản", style="B")
    rows = [[i, r.get("ngay_mua", ""), r["ma"], r["ten"], r["dac_diem"], r["dvt"], qty(r["sl"]),
             money(r["gia_mua"]), money(r["con_lai"])] for i, r in enumerate(k.items.to_dict("records"), start=1)]
    sl = pd.to_numeric(k.items["sl"], errors="coerce").fillna(0).sum()
    _pdf_table(pdf, ["STT", "Ngày mua", "Mã tài sản", "Tên tài sản", "Đặc điểm", "ĐVT", "SL", "Nguyên giá",
                     "Giá còn lại"], rows, [11, 18, 32, 38, 60, 13, 11, 28, 28], "CCCLLCCRR",
               total=["", "", "", "Tổng cộng", "", "", qty(sl),
                      money(pd.to_numeric(k.items["gia_mua"], errors="coerce").sum()),
                      money(pd.to_numeric(k.items["con_lai"], errors="coerce").sum())])
    _line(pdf, "Kết quả thanh lý tài sản", style="B")
    _line(pdf, f"Chi phí thanh lý tài sản: {k.chi_phi or 'không phát sinh'}.")
    _line(pdf, k.thu_hoi_text())
    _line(pdf, k.ghi_giam_text())
    _pdf_signatures(pdf, [("NGƯỜI LẬP BÁO CÁO", k.nguoi_lap), ("ĐƠN VỊ PHỤ TRÁCH KẾ TOÁN", k.ke_toan),
                          ("ĐƠN VỊ QUẢN LÝ TÀI SẢN/\nHỘI ĐỒNG THANH LÝ TÀI SẢN", k.quan_ly),
                          ("CẤP THẨM QUYỀN", k.tham_quyen)])
    return bytes(pdf.output())
