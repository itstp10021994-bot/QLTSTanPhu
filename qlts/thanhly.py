"""Thanh lý tài sản: Kế hoạch thanh lý (mẫu HC/QT–02/M05) và Tờ trình phê duyệt thanh lý.

Kế hoạch: Excel (A4 ngang) + PDF. Tờ trình: Word (điền vào file mẫu của trường) + PDF. Font Times New Roman.
"""

from __future__ import annotations

import copy
import io
from dataclasses import dataclass, field
from datetime import date

import pandas as pd

from .bienban import FONTS, LOGO, ROOT, _set_text
from .config import app_setting

TT_TEMPLATE = ROOT / "templates" / "to_trinh_thanh_ly.docx"

# Kế hoạch thanh lý: (khóa, tiêu đề, độ rộng Excel)
KH_COLUMNS = [
    ("stt", "Stt", 5), ("ngay_mua", "Ngày mua", 10), ("ma", "Mã tài sản", 16), ("ten", "Tên tài sản", 18),
    ("dac_diem", "Đặc điểm", 22), ("dvt", "Đvt", 6), ("sl", "SL", 6), ("gia_mua", "Giá mua mới", 14),
    ("con_lai", "Giá trị còn lại", 13), ("don_vi", "Đơn vị\nsử dụng", 14), ("tinh_trang", "Tình trạng", 19),
    ("du_kien", "Dự kiến\nthời gian thanh lý", 14), ("ghi_chu", "Ghi chú", 11),
]
# Tờ trình: bảng danh mục
TT_COLUMNS = ["stt", "ngay_mua", "ma", "ten", "dac_diem", "dvt", "sl", "don_vi", "tinh_trang", "ghi_chu"]
TT_HEADERS = ["Stt", "Ngày mua", "Mã tài sản", "Tên tài sản", "Đặc điểm", "Đvt", "SL", "Đơn vị\nsử dụng",
              "Tình trạng", "Ghi chú"]
MONEY = {"gia_mua", "con_lai"}
CENTER = {"stt", "ngay_mua", "ma", "dvt", "sl", "don_vi", "du_kien", "ghi_chu"}
CAN_CU = [
    "Căn cứ Quy chế quản thanh lý tài sản mã số QC-11 hiệu lực ngày 12/06/2024;",
    "Căn cứ Quy trình quản lý tài sản mã số HC/QT–02/M05 hiệu lực ngày 25/11/2025;",
    "Căn cứ tình hình thực tế hiện nay;",
]


def school_year(d: date) -> str:
    start = d.year if d.month >= 8 else d.year - 1
    return f"{start} - {start + 1}"


def kh_code() -> str:
    return "\n".join([
        f"Mã số: {app_setting('thanhly_ma_so', 'HC/QT–02/M05')}",
        f"Lần ban hành: {app_setting('thanhly_lan_ban_hanh', '08')}",
        f"Hiệu lực: {app_setting('thanhly_hieu_luc', '25/11/2025')}",
    ])


def money(v) -> str:
    try:
        return f"{float(v):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return str(v or "")


def qty(v) -> str:
    try:
        return f"{float(v):g}".replace(".", ",")
    except (TypeError, ValueError):
        return str(v or "")


def cell_text(key: str, value) -> str:
    if key in MONEY:
        return money(value)
    if key == "sl":
        return qty(value)
    return "" if value is None or (isinstance(value, float) and pd.isna(value)) else str(value)


def items_from(df: pd.DataFrame, room_names: dict[str, str], du_kien: str, ghi_chu: str) -> pd.DataFrame:
    """Thiết bị (từ SharePoint) -> các dòng biểu mẫu thanh lý."""
    sl = pd.to_numeric(df["SL"], errors="coerce").fillna(0)
    return pd.DataFrame({
        "id": df["id"].values,
        "ngay_mua": [str(v)[:4] if v else "" for v in df["NgayMua"]],
        "ma": df["MaChiTiet"].values,
        "ten": df["TenThietBi"].where(df["TenThietBi"] != "", df["ChiTiet"]).values,
        "dac_diem": df["DacDiem"].values,
        "dvt": df["DVT"].replace("", "Cái").values,
        "sl": sl.where(sl > 0, 1).values,
        "gia_mua": pd.to_numeric(df["GiaTri"], errors="coerce").fillna(0).values,
        "con_lai": 0.0,
        "don_vi": [room_names.get(r, r) for r in df["NoiSuDung"]],
        "tinh_trang": df["TinhTrang"].values,
        "du_kien": du_kien,
        "ghi_chu": ghi_chu,
    })


@dataclass
class ThanhLy:
    items: pd.DataFrame
    ngay: date = field(default_factory=date.today)
    so: str = ""
    trich_yeu: str = ""
    kinh_gui: str = "Ban giám hiệu Trường TH, THCS và THPT Tân Phú"
    noi_dung: str = ""
    hinh_thuc: str = ""
    bao_cao: str = ""
    trien_khai: date | None = None
    tieu_de_phu: str = ""
    ky_quan_ly: str = ""
    ky_ke_toan: str = ""
    ky_phe_duyet: str = ""

    @property
    def total(self) -> float:
        return float(pd.to_numeric(self.items["sl"], errors="coerce").fillna(0).sum()) if not self.items.empty else 0

    def date_line(self) -> str:
        return f"Tp Hồ Chí Minh, ngày {self.ngay.day} tháng {self.ngay.month} năm {self.ngay.year}"

    def so_line(self) -> str:
        return f"Số: {self.so or '....'} /TTr-TP"

    def trien_khai_text(self) -> str:
        d = self.trien_khai or self.ngay
        return f" Kể từ ngày {d.day}/{d.month}/{d.year}./."


# ---------------------------------------------------------------------------
# Kế hoạch thanh lý – Excel
# ---------------------------------------------------------------------------
def kh_xlsx(tl: ThanhLy) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter as col

    wb = Workbook()
    ws = wb.active
    ws.title = "KeHoachThanhLy"
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    f = lambda size=12, bold=False: Font(name="Times New Roman", size=size, bold=bold)  # noqa: E731
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    n = len(KH_COLUMNS)
    last = col(n)
    for i, (_, _, w) in enumerate(KH_COLUMNS, start=1):
        ws.column_dimensions[col(i)].width = w

    def put(rng, text, size=12, bold=False, align=center):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value, c.font, c.alignment = text, f(size, bold), align

    put("A1:C1", None)
    put("D1:K1", "KẾ HOẠCH THANH LÝ", 14, True)
    put(f"L1:{last}1", kh_code(), 10)
    ws.row_dimensions[1].height = 48
    for i in range(1, n + 1):
        ws.cell(1, i).border = box
    if LOGO.exists():
        img = XLImage(str(LOGO))
        img.height, img.width = 56, 56 * 159 / 99
        ws.add_image(img, "A1")
    put(f"A2:{last}2", tl.tieu_de_phu, 12, False)
    ws.row_dimensions[2].height = 22

    r = 3
    for i, (_, label, _) in enumerate(KH_COLUMNS, start=1):
        c = ws.cell(r, i, label)
        c.font, c.alignment, c.border = f(12, True), center, box
    ws.row_dimensions[r].height = 32
    for k, row in enumerate(tl.items.to_dict("records"), start=1):
        r += 1
        for i, (key, _, _) in enumerate(KH_COLUMNS, start=1):
            v = k if key == "stt" else row.get(key, "")
            if key in MONEY or key == "sl":
                v = float(v or 0)
            c = ws.cell(r, i, v)
            c.font, c.border = f(12), box
            c.alignment = center if key in CENTER else left
            if key in MONEY:
                c.number_format = "#,##0"
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "Tổng cộng").font = f(12, True)
    ws.cell(r, 1).alignment = center
    ws.cell(r, 7, tl.total).font = f(12, True)
    ws.cell(r, 7).alignment = center
    for i in range(1, n + 1):
        ws.cell(r, i).border = box

    r += 3
    put(f"A{r}:F{r}", "Đơn vị quản lý tài sản", 11, True)
    put(f"G{r}:J{r}", "Đơn vị phụ trách kế toán", 11, True)
    put(f"K{r}:{last}{r}", "Phê duyệt", 11, True)
    r += 5
    put(f"A{r}:F{r}", tl.ky_quan_ly, 11)
    put(f"G{r}:J{r}", tl.ky_ke_toan, 11)
    put(f"K{r}:{last}{r}", tl.ky_phe_duyet, 11)

    ws.print_title_rows = "3:3"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.print_options.horizontalCentered = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF dùng chung
# ---------------------------------------------------------------------------
def _pdf(orientation: str):
    from fpdf import FPDF

    pdf = FPDF(orientation=orientation, format="A4", unit="mm")
    pdf.add_font("TNR", "", str(FONTS / "LiberationSerif-Regular.ttf"))
    pdf.add_font("TNR", "B", str(FONTS / "LiberationSerif-Bold.ttf"))
    pdf.add_font("TNR", "I", str(FONTS / "LiberationSerif-Italic.ttf"))
    return pdf


def _table(pdf, keys: list[str], headers: list[str], widths: list[float], tl: ThanhLy, size: float,
           total_span: int) -> None:
    from fpdf.fonts import FontFace

    width = pdf.w - pdf.l_margin - pdf.r_margin
    margin = pdf.l_margin
    pdf.set_font("TNR", "", size)
    align = tuple("C" if k in CENTER else ("R" if k in MONEY else "L") for k in keys)
    with pdf.table(width=width, align="LEFT", col_widths=widths, text_align=align, line_height=size * 0.5,
                   headings_style=FontFace(emphasis="BOLD", size_pt=size), padding=1) as table:
        row = table.row()
        for h in headers:
            row.cell(h, align="C")
        for k, item in enumerate(tl.items.to_dict("records"), start=1):
            row = table.row()
            for key in keys:
                row.cell(str(k) if key == "stt" else cell_text(key, item.get(key, "")))
        row = table.row(style=FontFace(emphasis="BOLD"))
        row.cell("Tổng cộng", colspan=total_span, align="C")
        row.cell(qty(tl.total), align="C")
        rest = len(keys) - total_span - 1
        if rest:
            row.cell("", colspan=rest)
    pdf.set_left_margin(margin)


def kh_pdf(tl: ThanhLy) -> bytes:
    pdf = _pdf("L")
    pdf.set_margins(8, 10, 8)
    pdf.set_auto_page_break(True, 12)
    pdf.add_page()
    width = pdf.w - 16
    widths = [w for _, _, w in KH_COLUMNS]
    scale = width / sum(widths)
    xs = [8 + sum(widths[:i]) * scale for i in range(len(widths) + 1)]
    y0, h0 = pdf.get_y(), 18
    pdf.rect(8, y0, width, h0)
    pdf.line(xs[3], y0, xs[3], y0 + h0)
    pdf.line(xs[11], y0, xs[11], y0 + h0)
    if LOGO.exists():
        lw = (h0 - 4) * 159 / 99
        pdf.image(str(LOGO), x=8 + (xs[3] - 8 - lw) / 2, y=y0 + 2, h=h0 - 4)
    pdf.set_font("TNR", "B", 14)
    pdf.set_xy(xs[3], y0)
    pdf.cell(xs[11] - xs[3], h0, "KẾ HOẠCH THANH LÝ", align="C")
    pdf.set_font("TNR", "", 9.5)
    pdf.set_xy(xs[11], y0 + 2.5)
    pdf.multi_cell(xs[-1] - xs[11], 4.5, kh_code(), align="C")
    pdf.set_y(y0 + h0 + 1)
    if tl.tieu_de_phu:
        pdf.set_font("TNR", "B", 12)
        pdf.cell(width, 8, tl.tieu_de_phu, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    _table(pdf, [k for k, _, _ in KH_COLUMNS], [h for _, h, _ in KH_COLUMNS], widths, tl, 10.5, 6)
    pdf.ln(5)
    if pdf.get_y() > pdf.h - pdf.b_margin - 34:  # giữ phần ký tên trên cùng một trang
        pdf.add_page()
    y = pdf.get_y()
    blocks = [(xs[0], xs[6], "Đơn vị quản lý tài sản", tl.ky_quan_ly),
              (xs[6], xs[10], "Đơn vị phụ trách kế toán", tl.ky_ke_toan),
              (xs[10], xs[-1], "Phê duyệt", tl.ky_phe_duyet)]
    for a, b, title, name in blocks:
        pdf.set_font("TNR", "B", 11.5)
        pdf.set_xy(a, y)
        pdf.cell(b - a, 6, title, align="C")
        pdf.set_font("TNR", "", 11.5)
        pdf.set_xy(a, y + 26)
        pdf.cell(b - a, 6, name, align="C")
    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Tờ trình – Word (điền vào file mẫu)
# ---------------------------------------------------------------------------
def _set_after_label(paragraph, text: str) -> None:
    """Giữ các run đậm đầu đoạn (nhãn), thay phần nội dung phía sau bằng ``text`` (định dạng thường)."""
    runs = paragraph.runs
    n = 0
    while n < len(runs) and runs[n].bold:
        n += 1
    plain = next((r for r in runs[n:] if not r.bold), None)
    proto = copy.deepcopy(plain._r) if plain is not None else None
    for r in runs[n:]:
        r._r.getparent().remove(r._r)
    if proto is not None:
        for t in proto.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
            proto.remove(t)
        paragraph._p.append(proto)
        from docx.text.run import Run
        Run(proto, paragraph).text = text
    else:
        run = paragraph.add_run(text)
        run.bold = False


def tt_docx(tl: ThanhLy) -> bytes:
    from docx import Document

    doc = Document(str(TT_TEMPLATE))
    head = doc.tables[0]
    for cell in (head.rows[-1].cells[0], head.rows[-1].cells[-1]):
        p = cell.paragraphs[0]
        if "Số" in p.text:
            _set_text(p, "                     " + tl.so_line())
        elif "ngày" in p.text:
            _set_text(p, tl.date_line())
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("Vv "):
            _set_text(p, tl.trich_yeu)
        elif text.startswith("Kính gửi"):
            _set_text(p, "Kính gửi: " + tl.kinh_gui)
        elif text.startswith("Trên cơ sở"):
            _set_text(p, "        " + tl.noi_dung)
        elif text.startswith("Hình thức thanh lý"):
            p.runs[0].text = "Hình thức thanh lý:"
            _set_after_label(p, " " + tl.hinh_thuc)
        elif text.startswith("Báo cáo kết quả thanh lý"):
            _set_after_label(p, " " + tl.bao_cao)
        elif text.startswith("Thời điểm triển khai"):
            _set_after_label(p, tl.trien_khai_text())
    tbl = doc.tables[1]
    proto = copy.deepcopy(tbl.rows[1]._tr)
    total_tr = tbl.rows[-1]._tr
    for row in list(tbl.rows)[1:-1]:
        tbl._tbl.remove(row._tr)
    for k, item in enumerate(tl.items.to_dict("records"), start=1):
        tr = copy.deepcopy(proto)
        total_tr.addprevious(tr)
        from docx.table import _Row
        cells = _Row(tr, tbl).cells
        for cell, key in zip(cells, TT_COLUMNS):
            _set_text(cell.paragraphs[0], str(k) if key == "stt" else cell_text(key, item.get(key, "")))
    _set_text(tbl.rows[-1].cells[6].paragraphs[0], qty(tl.total))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def tt_pdf(tl: ThanhLy) -> bytes:
    pdf = _pdf("P")
    pdf.set_margins(22, 14, 16)
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    half = width * 0.46
    y = pdf.get_y()
    pdf.set_font("TNR", "", 11.5)
    pdf.multi_cell(half, 5.5, "CÔNG TY CỔ PHẦN GIÁO DỤC\nTHÀNH THÀNH CÔNG", align="C")
    pdf.set_font("TNR", "B", 11)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(half, 5.5, "TRƯỜNG TH, THCS VÀ THPT TÂN PHÚ", align="C")
    bottom = pdf.get_y()
    pdf.set_font("TNR", "B", 11)
    pdf.set_xy(pdf.l_margin + half, y)
    pdf.multi_cell(width - half, 5.5, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc", align="C")
    pdf.set_y(max(bottom, pdf.get_y()) + 4)
    pdf.set_font("TNR", "", 13)
    pdf.cell(half, 6, tl.so_line(), align="C")
    pdf.set_font("TNR", "I", 13)
    pdf.cell(width - half, 6, tl.date_line(), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    pdf.set_font("TNR", "B", 14)
    pdf.cell(width, 7, "TỜ TRÌNH", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("TNR", "", 13)
    pdf.multi_cell(width, 6.5, tl.trich_yeu, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    def para(label: str, text: str = "", style: str = "", indent: bool = False):
        pdf.set_x(pdf.l_margin)
        if label:
            pdf.set_font("TNR", "B", 13)
            pdf.write(6.5, label)
        pdf.set_font("TNR", style, 13)
        pdf.write(6.5, ("        " if indent else "") + text)
        pdf.ln(6.5)
        pdf.ln(1)

    para("Kính gửi: ", tl.kinh_gui, "B")
    pdf.ln(2)
    for line in CAN_CU:
        para("", line, "I")
    para("", tl.noi_dung, indent=True)
    para("Danh mục tài sản thanh lý:")
    widths = [8, 11, 25, 20, 30, 9, 8, 19, 22, 17]
    _table(pdf, TT_COLUMNS, TT_HEADERS, widths, tl, 10.5, 6)
    pdf.ln(4)
    para("Hình thức thanh lý: ", tl.hinh_thuc)
    para("Báo cáo kết quả thanh lý: ", tl.bao_cao)
    para("Thời điểm triển khai thực hiện:", tl.trien_khai_text())
    pdf.ln(4)
    para("", "Trân trọng kính trình Ban giám hiệu xem xét và thuận duyệt./.", "I")
    pdf.ln(8)
    if pdf.get_y() > pdf.h - 55:
        pdf.add_page()
    for label in ("ĐƠN VỊ TRÌNH: ", "ĐƠN VỊ THAM MƯU: ", "PHÊ DUYỆT: "):
        pdf.set_font("TNR", "B", 13)
        pdf.set_x(pdf.l_margin)
        pdf.cell(width, 6.5, label + "." * (100 - len(label) * 2), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(9)
    return bytes(pdf.output())
