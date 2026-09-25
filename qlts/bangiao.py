"""Biên bản bàn giao tài sản theo đơn vị sử dụng (mẫu HC/QT–02/M03): gom nhóm, xuất Excel / PDF.

Font Times New Roman (PDF dùng Liberation Serif – cùng kích thước chữ với Times New Roman).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date

import pandas as pd

from .bienban import FONTS, LOGO, _sheet_name
from .config import app_setting

# (khóa, tiêu đề cột, độ rộng Excel)
COLUMNS = [
    ("stt", "STT", 6),
    ("nhom", "Nhóm tài sản", 29),
    ("ma", "Mã tài sản", 13),
    ("ten", "Tên tài sản", 24),
    ("nguoi", "Nhân sự sử dụng", 22),
    ("sl", "Số lượng", 9),
    ("tinh_trang", "Tình trạng", 13),
    ("vi_tri", "Vị trí tài sản", 15),
    ("quy_cach", "Quy cách", 31),
]
LABELS = {k: label for k, label, _ in COLUMNS}


NO_PLACE = "(Chưa có vị trí)"


def school_year(d: date) -> str:
    start = d.year if d.month >= 8 else d.year - 1
    return f"{start}-{start + 1}"


def header_code() -> str:
    return "\n".join([
        f"Mã số: {app_setting('bangiao_ma_so', 'HC/QT–02/M03')}",
        f"Lần ban hành: {app_setting('bangiao_lan_ban_hanh', '08')}",
        f"Hiệu lực: {app_setting('bangiao_hieu_luc', '25/11/2025')}",
    ])


def place() -> str:
    return str(app_setting("bangiao_dia_diem", "Trường TH, THCS và THPT Tân Phú"))


def group_assets(df: pd.DataFrame) -> pd.DataFrame:
    """Gom thiết bị cùng nơi sử dụng + mã tài sản + quy cách + người sử dụng + tình trạng; cộng số lượng.

    Quy cách = Chi tiết (nếu Chi tiết trống hoặc trùng tên thiết bị thì lấy Đặc điểm).
    """
    cols = ["vi_tri", "nhom", "ma", "ten", "quy_cach", "nguoi", "tinh_trang", "sl"]
    if df.empty:
        return pd.DataFrame(columns=cols)
    w = pd.DataFrame({
        "vi_tri": df["NoiSuDung"].str.strip().replace("", NO_PLACE),
        "nhom": df["NhomThietBi"].str.strip(),
        "ma": df["MaTaiSan"].str.strip(),
        "ten": df["TenThietBi"].where(df["TenThietBi"].str.strip() != "", df["ChiTiet"]).str.strip(),
        "nguoi": df["NguoiSuDung"].str.strip(),
        "tinh_trang": df["TinhTrang"].str.strip(),
    })
    same = df["ChiTiet"].str.strip().str.lower() == df["TenThietBi"].str.strip().str.lower()
    w["quy_cach"] = df["ChiTiet"].where((df["ChiTiet"].str.strip() != "") & ~same, df["DacDiem"]).str.strip()
    sl = pd.to_numeric(df["SL"], errors="coerce").fillna(0)
    w["sl"] = sl.where(sl > 0, 1)
    keys = ["vi_tri", "nhom", "ma", "ten", "quy_cach", "nguoi", "tinh_trang"]
    out = w.groupby(keys, as_index=False, sort=False)["sl"].sum()
    return out.sort_values(["vi_tri", "nhom", "ma", "quy_cach"], kind="stable").reset_index(drop=True)[cols]


@dataclass
class BanGiao:
    don_vi: str  # tên hiển thị của đơn vị sử dụng (phòng)
    dai_dien: str
    dai_dien_cv: str
    ban: list[tuple[str, str]]  # Ban bàn giao: (họ tên, chức vụ)
    rows: pd.DataFrame  # kết quả group_assets của phòng
    ngay: date = field(default_factory=date.today)
    nam_hoc: str = ""
    sheet: str = "Ban giao"

    @property
    def title(self) -> str:
        return f"BIÊN BẢN BÀN GIAO TÀI SẢN\nNĂM {self.nam_hoc or school_year(self.ngay)}"

    def date_line(self) -> str:
        return (f"Hôm nay, ngày {self.ngay.day:02d} tháng {self.ngay.month:02d} năm {self.ngay.year}, "
                f"tại {place()} chúng tôi gồm:")


INTRO = ("Chúng tôi cùng thống nhất kiểm kê số lượng và đánh giá % sử dụng còn lại của từng tài sản "
         "và công cụ dụng cụ như sau:")


def _fmt_qty(value) -> str:
    return f"{float(value):g}".replace(".", ",")


def _row_values(i: int, r) -> list:
    return [i, r.nhom, r.ma, r.ten, r.nguoi, r.sl, r.tinh_trang, r.vi_tri, r.quy_cach]


# ---------------------------------------------------------------------------
# Excel – mỗi đơn vị một sheet, A4 ngang, Times New Roman
# ---------------------------------------------------------------------------
def to_xlsx(docs: list[BanGiao]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def font(size=12, bold=False, italic=False):
        return Font(name="Times New Roman", size=size, bold=bold, italic=italic)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    last = get_column_letter(len(COLUMNS))  # I

    for bb in docs:
        ws = wb.create_sheet(_sheet_name(bb.sheet, used))
        for n, (_, _, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(n)].width = width

        def put(cell_range: str, text, size=12, bold=False, align=left, height=None):
            ws.merge_cells(cell_range)
            first = cell_range.split(":")[0]
            c = ws[first]
            c.value, c.font, c.alignment = text, font(size, bold), align
            if height:
                ws.row_dimensions[c.row].height = height

        # Dòng tiêu đề: logo | tên biên bản | mã số
        put("A1:C1", None, align=center, height=52)
        put("D1:G1", bb.title, size=14, bold=True, align=center)
        put(f"H1:{last}1", header_code(), size=10, align=center)
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(1, col).border = box
        if LOGO.exists():
            img = XLImage(str(LOGO))
            img.height, img.width = 58, 58 * 159 / 99
            ws.add_image(img, "B1")

        r = 3
        put(f"A{r}:{last}{r}", bb.date_line()); r += 1
        put(f"A{r}:{last}{r}", "I. ĐẠI DIỆN BAN BÀN GIAO TÀI SẢN", bold=True); r += 1
        for n, (name, cv) in enumerate(bb.ban, start=1):
            put(f"A{r}:B{r}", f"{n}. Ông/Bà:")
            put(f"C{r}:D{r}", name)
            ws[f"E{r}"].value, ws[f"E{r}"].font, ws[f"E{r}"].alignment = "Chức vụ:", font(), center
            put(f"F{r}:{last}{r}", cv)
            r += 1
        put(f"A{r}:{last}{r}", f"II. ĐƠN VỊ SỬ DỤNG TÀI SẢN: {bb.don_vi}", bold=True); r += 1
        put(f"A{r}:B{r}", "Đại diện đơn vị:")
        put(f"C{r}:D{r}", bb.dai_dien)
        ws[f"E{r}"].value, ws[f"E{r}"].font, ws[f"E{r}"].alignment = "Chức vụ:", font(), center
        put(f"F{r}:{last}{r}", bb.dai_dien_cv); r += 1
        put(f"A{r}:{last}{r}", INTRO); r += 2

        head = r
        for n, (_, label, _) in enumerate(COLUMNS, start=1):
            c = ws.cell(r, n, label)
            c.font, c.alignment, c.border = font(11, bold=True), center, box
        ws.row_dimensions[r].height = 30
        r += 1
        for i, row in enumerate(bb.rows.itertuples(), start=1):
            for n, v in enumerate(_row_values(i, row), start=1):
                c = ws.cell(r, n, v)
                c.font, c.border = font(11), box
                c.alignment = center if COLUMNS[n - 1][0] in ("stt", "ma", "sl", "tinh_trang", "vi_tri") else left
            r += 1
        c = ws.cell(r, 1, "Tổng cộng")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c.font, c.alignment = font(11, bold=True), center
        tot = ws.cell(r, 6, float(bb.rows["sl"].sum()) if not bb.rows.empty else 0)
        tot.font, tot.alignment = font(11, bold=True), center
        for n in range(1, len(COLUMNS) + 1):
            ws.cell(r, n).border = box
        r += 2
        put(f"A{r}:D{r}", "BAN BÀN GIAO", bold=True, align=center)
        put(f"F{r}:{last}{r}", "ĐƠN VỊ SỬ DỤNG", bold=True, align=center)
        r += 5
        put(f"A{r}:D{r}", "\n".join(n for n, _ in bb.ban if n), bold=True, align=center,
            height=16 * max(1, len([n for n, _ in bb.ban if n])))
        put(f"F{r}:{last}{r}", bb.dai_dien, bold=True, align=center)

        ws.print_title_rows = f"{head}:{head}"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left, ws.page_margins.right, ws.page_margins.top = 0.4, 0.25, 0.4
        ws.print_options.horizontalCentered = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF – mỗi đơn vị bắt đầu một trang mới, A4 ngang
# ---------------------------------------------------------------------------
def to_pdf(docs: list[BanGiao]) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    pdf = FPDF(orientation="L", format="A4", unit="mm")
    pdf.set_margins(10, 10, 8)
    pdf.set_auto_page_break(True, 12)
    pdf.add_font("TNR", "", str(FONTS / "LiberationSerif-Regular.ttf"))
    pdf.add_font("TNR", "B", str(FONTS / "LiberationSerif-Bold.ttf"))
    pdf.add_font("TNR", "I", str(FONTS / "LiberationSerif-Italic.ttf"))
    width = pdf.w - pdf.l_margin - pdf.r_margin
    widths = [w for _, _, w in COLUMNS]
    scale = width / sum(widths)

    def x_of(col: int) -> float:
        return pdf.l_margin + sum(widths[:col]) * scale

    def line(text: str, style: str = "", h: float = 6):
        pdf.set_font("TNR", style, 12)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(width, h, text, new_x="LMARGIN", new_y="NEXT")

    def person(label: str, name: str, cv: str):
        pdf.set_font("TNR", "", 12)
        y = pdf.get_y()
        pdf.set_xy(x_of(0), y)
        pdf.cell(x_of(2) - x_of(0), 6, label)
        pdf.set_xy(x_of(2), y)
        pdf.cell(x_of(4) - x_of(2), 6, name)
        pdf.set_xy(x_of(4), y)
        pdf.cell(x_of(5) - x_of(4), 6, "Chức vụ:", align="C")
        pdf.set_xy(x_of(5), y)
        pdf.cell(width - (x_of(5) - pdf.l_margin), 6, cv, new_x="LMARGIN", new_y="NEXT")

    for bb in docs:
        pdf.add_page()
        y0, h0 = pdf.get_y(), 18
        a, b = x_of(3), x_of(7)
        pdf.set_line_width(0.3)
        pdf.rect(pdf.l_margin, y0, width, h0)
        pdf.line(a, y0, a, y0 + h0)
        pdf.line(b, y0, b, y0 + h0)
        if LOGO.exists():
            logo_w = (h0 - 4) * 159 / 99
            pdf.image(str(LOGO), x=pdf.l_margin + (a - pdf.l_margin - logo_w) / 2, y=y0 + 2, h=h0 - 4)
        pdf.set_font("TNR", "B", 14)
        pdf.set_xy(a, y0 + 3.5)
        pdf.multi_cell(b - a, 6.5, bb.title, align="C")
        pdf.set_font("TNR", "", 10)
        pdf.set_xy(b, y0 + 2.5)
        pdf.multi_cell(pdf.l_margin + width - b, 5, header_code(), align="C")
        pdf.set_y(y0 + h0 + 3)

        line(bb.date_line())
        line("I. ĐẠI DIỆN BAN BÀN GIAO TÀI SẢN", "B")
        for n, (name, cv) in enumerate(bb.ban, start=1):
            person(f"{n}. Ông/Bà:", name, cv)
        line(f"II. ĐƠN VỊ SỬ DỤNG TÀI SẢN: {bb.don_vi}", "B")
        person("Đại diện đơn vị:", bb.dai_dien, bb.dai_dien_cv)
        line(INTRO)
        pdf.ln(2)

        pdf.set_font("TNR", "", 10)
        head = FontFace(emphasis="BOLD", size_pt=10)
        align = tuple("C" if k in ("stt", "ma", "sl", "tinh_trang", "vi_tri") else "L" for k, _, _ in COLUMNS)
        margin = pdf.l_margin
        with pdf.table(width=width, align="LEFT", col_widths=widths, text_align=align, headings_style=head,
                       line_height=5, padding=1) as table:
            row = table.row()
            for _, label, _ in COLUMNS:
                row.cell(label, align="C")
            for i, r in enumerate(bb.rows.itertuples(), start=1):
                row = table.row()
                for k, v in zip([c[0] for c in COLUMNS], _row_values(i, r)):
                    row.cell(_fmt_qty(v) if k == "sl" else str(v))
            row = table.row(style=FontFace(emphasis="BOLD"))
            row.cell("Tổng cộng", colspan=5, align="C")
            row.cell(_fmt_qty(bb.rows["sl"].sum() if not bb.rows.empty else 0), align="C")
            row.cell("", colspan=3)
        pdf.set_left_margin(margin)
        pdf.ln(4)

        names = [n for n, _ in bb.ban if n]
        if pdf.get_y() + 26 + 6 * len(names) > pdf.h - pdf.b_margin:  # giữ phần ký tên trên cùng một trang
            pdf.add_page()
        y = pdf.get_y()
        pdf.set_font("TNR", "B", 12)
        pdf.set_xy(x_of(0), y)
        pdf.cell(x_of(4) - x_of(0), 6, "BAN BÀN GIAO", align="C")
        pdf.set_xy(x_of(5), y)
        pdf.cell(pdf.l_margin + width - x_of(5), 6, "ĐƠN VỊ SỬ DỤNG", align="C")
        pdf.set_xy(x_of(0), y + 20)
        pdf.multi_cell(x_of(4) - x_of(0), 6, "\n".join(names), align="C")
        pdf.set_xy(x_of(5), y + 20)
        pdf.multi_cell(pdf.l_margin + width - x_of(5), 6, bb.dai_dien, align="C")
    return bytes(pdf.output())


def to_html(bb: BanGiao) -> str:
    """Xem trước biên bản trên màn hình (Times New Roman), giống bản in."""
    import base64
    from html import escape as e

    logo = (f"<img src='data:image/jpeg;base64,{base64.b64encode(LOGO.read_bytes()).decode()}' height='56'>"
            if LOGO.exists() else "")

    ban = "".join(f"<tr><td>{n}. Ông/Bà:</td><td>{e(name)}</td><td>Chức vụ:</td><td>{e(cv)}</td></tr>"
                  for n, (name, cv) in enumerate(bb.ban, start=1))
    head = "".join(f"<th>{e(label)}</th>" for _, label, _ in COLUMNS)
    body = "".join(
        "<tr>" + "".join(
            f"<td class='{'c' if k in ('stt', 'ma', 'sl', 'tinh_trang', 'vi_tri') else ''}'>"
            f"{e(_fmt_qty(v) if k == 'sl' else str(v))}</td>"
            for (k, _, _), v in zip(COLUMNS, _row_values(i, r))) + "</tr>"
        for i, r in enumerate(bb.rows.itertuples(), start=1))
    total = _fmt_qty(bb.rows["sl"].sum() if not bb.rows.empty else 0)
    return f"""
<div class="bgp">
<style>
.bgp {{font-family:'Times New Roman',Times,serif; font-size:15px; color:inherit}}
.bgp table {{border-collapse:collapse; width:100%}}
.bgp .hd td {{border:1px solid #888; text-align:center; padding:6px}}
.bgp .hd .t {{font-weight:bold; font-size:18px}}
.bgp .hd .m {{font-size:12px}}
.bgp table.ps {{width:auto}}
.bgp .ps td {{padding:1px 28px 1px 0; border:none}}
.bgp .grid th, .bgp .grid td {{border:1px solid #888; padding:3px 5px; font-size:14px}}
.bgp .grid th {{text-align:center}}
.bgp .grid .c {{text-align:center}}
.bgp b {{display:block; margin-top:6px}}
</style>
<table class="hd"><tr><td style="width:30%">{logo}</td>
<td class="t">{e(bb.title).replace(chr(10), "<br>")}</td>
<td class="m" style="width:25%">{e(header_code()).replace(chr(10), "<br>")}</td></tr></table>
<p style="margin:8px 0 0">{e(bb.date_line())}</p>
<b>I. ĐẠI DIỆN BAN BÀN GIAO TÀI SẢN</b>
<table class="ps">{ban}</table>
<b>II. ĐƠN VỊ SỬ DỤNG TÀI SẢN: {e(bb.don_vi)}</b>
<table class="ps"><tr><td>Đại diện đơn vị:</td><td>{e(bb.dai_dien)}</td><td>Chức vụ:</td>
<td>{e(bb.dai_dien_cv)}</td></tr></table>
<p style="margin:2px 0 6px">{e(INTRO)}</p>
<table class="grid"><tr>{head}</tr>{body}
<tr><td colspan="5" class="c"><strong>Tổng cộng</strong></td><td class="c"><strong>{total}</strong></td>
<td colspan="3"></td></tr></table>
</div>"""
