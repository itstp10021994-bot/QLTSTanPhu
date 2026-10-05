"""Báo cáo tuyển sinh qua các năm từ 2 list SharePoint có sẵn: "Nhập học" và "Data tuyển sinh".

Hai list do trường tự tạo nên tên cột không cố định: app dò cột theo tên hiển thị (``guess_columns``),
người dùng chỉnh lại được trên trang. Dữ liệu quận/huyện, phường/xã thường ghi không thống nhất
("Q.Tân Bình", "tân phú", "H. Bình Chánh"...) nên được chuẩn hóa trước khi thống kê; nếu list không có
cột quận/huyện riêng thì tách từ cột Địa chỉ. Báo cáo chỉ hiển thị số tổng hợp, không hiển thị hồ sơ học sinh.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date

import pandas as pd

UNKNOWN = "(Chưa rõ)"
TOTAL = "Tổng cộng"
DATA, NH = "Data tuyển sinh", "Nhập học"

# khóa: (nhãn, từ khóa tên cột – chữ thường, bỏ dấu, viết liền; ưu tiên theo thứ tự)
FIELDS: dict[str, tuple[str, list[str]]] = {
    "nam": ("Năm học / năm tuyển sinh", ["namhoc", "nienkhoa", "namtuyensinh", "khoatuyensinh", "namnhaphoc",
                                          "nam", "khoa", "schoolyear"]),
    "ngay": ("Ngày (suy ra năm học khi không có cột năm)", ["ngaynhaphoc", "ngaydangky", "ngaynophoso", "ngaynop",
                                                          "ngaytuvan", "ngaytao", "created", "ngay"]),
    "tinh": ("Tỉnh / thành phố", ["tinhthanhpho", "tinhthanh", "tinhtp", "tinh", "thanhpho"]),
    "quan": ("Quận / huyện", ["quanhuyen", "quanhuyenthixa", "quanhuyentp", "quan", "huyen", "district"]),
    "phuong": ("Phường / xã", ["phuongxa", "phuongxathitran", "phuong", "xa", "ward"]),
    "diachi": ("Địa chỉ (tách quận/huyện, phường/xã khi không có cột riêng)",
               ["diachi", "diachithuongtru", "diachitamtru", "diachinha", "diachilienhe", "noio", "address"]),
    "khoi": ("Khối / lớp", ["khoidangky", "khoi", "khoilop", "lopdangky", "lop", "lophoc", "grade"]),
    "cap": ("Cấp học", ["caphoc", "cap", "bachoc"]),
    "gioitinh": ("Giới tính", ["gioitinh", "phai"]),
    "nguon": ("Nguồn / kênh", ["nguon", "nguonthongtin", "kenh", "kenhtuyensinh", "nguontuyensinh"]),
    "trangthai": ("Trạng thái", ["trangthai", "tinhtrang", "ketqua"]),
}
DIMS = {"QuanHuyen": "Quận/Huyện", "PhuongXa": "Phường/Xã", "Tinh": "Tỉnh/TP", "Khoi": "Khối", "Cap": "Cấp học",
        "GioiTinh": "Giới tính", "Nguon": "Nguồn", "TrangThai": "Trạng thái"}


# ---------------------------------------------------------------------------
# Chuẩn hóa chữ
# ---------------------------------------------------------------------------
def _nfc(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value))).strip(" ,;.-\t")


def plain(text) -> str:
    """Chữ thường, bỏ dấu tiếng Việt (để so khớp)."""
    text = unicodedata.normalize("NFD", _nfc(text).lower()).replace("đ", "d")
    return "".join(c for c in text if unicodedata.category(c) != "Mn")


def _key(text) -> str:
    return re.sub(r"[^a-z0-9]", "", plain(text))


def _title(text: str) -> str:
    return " ".join(w[:1].upper() + w[1:].lower() for w in _nfc(text).split())


def guess_columns(columns: list[str]) -> dict[str, str | None]:
    """Đoán cột cho từng trường theo tên hiển thị; mỗi cột chỉ dùng một lần."""
    keys = {c: _key(c) for c in columns}
    used: set[str] = set()
    result: dict[str, str | None] = {}
    for field, (_, words) in FIELDS.items():
        found = None
        for w in words:  # khớp đúng
            found = next((c for c in columns if keys[c] == w and c not in used), None)
            if found:
                break
        if not found:  # bắt đầu bằng từ khóa (từ khóa đủ dài để tránh nhầm)
            for w in (w for w in words if len(w) >= 4):
                found = next((c for c in columns if keys[c].startswith(w) and c not in used), None)
                if found:
                    break
        result[field] = found
        if found:
            used.add(found)
    return result


# ---- Năm học ----
def year_start(label: str) -> int:
    m = re.search(r"\d{4}", label or "")
    return int(m.group()) if m else 9999


def school_year(value) -> str:
    """'2023-2024', '2023 - 2024', '23-24', 2023 -> '2023-2024' (số đơn lẻ = năm bắt đầu năm học)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (int, float)):
        y = int(value)
        return f"{y}-{y + 1}" if 1990 < y < 2100 else ""
    text = _nfc(value)
    years = [int(y) for y in re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", text)]
    if len(years) >= 2:
        return f"{years[0]}-{years[1]}"
    if len(years) == 1:
        return f"{years[0]}-{years[0] + 1}"
    m = re.fullmatch(r"(\d{2})\s*[-/–]\s*(\d{2})", text)
    if m:
        return f"20{m.group(1)}-20{m.group(2)}"
    return ""


def year_from_date(series: pd.Series) -> pd.Series:
    """Năm học suy ra từ ngày: từ tháng 9 trở đi tính cho năm học kế tiếp."""
    d = pd.to_datetime(series, errors="coerce", utc=True, format="mixed", dayfirst=False)
    start = d.dt.year + (d.dt.month >= 9).astype("Int64")
    return start.map(lambda y: f"{int(y)}-{int(y) + 1}" if pd.notna(y) else "")


# ---- Địa danh ----
HUYEN = {"binh chanh", "cu chi", "hoc mon", "nha be", "can gio"}
QUAN_TEN = {"tan phu", "tan binh", "binh tan", "binh thanh", "go vap", "phu nhuan"}
THANH_PHO = {"thu duc", "thuan an", "di an", "thu dau mot", "bien hoa", "tan uyen", "ben cat", "long khanh",
             "vung tau", "ba ria", "tan an"}
TINH_LON = {"ho chi minh": "TP. Hồ Chí Minh", "hcm": "TP. Hồ Chí Minh", "tphcm": "TP. Hồ Chí Minh",
            "sai gon": "TP. Hồ Chí Minh", "ha noi": "TP. Hà Nội", "da nang": "TP. Đà Nẵng",
            "can tho": "TP. Cần Thơ", "hai phong": "TP. Hải Phòng", "hue": "TP. Huế"}
_RE_QUAN_SO = re.compile(r"^(?:quận|quan|q)\s*\.?\s*0*(\d{1,2})$", re.I)
_PREFIXES = [  # (biểu thức, tiền tố chuẩn)
    (re.compile(r"^(?:thành phố|thanh pho|tp)\s*\.?\s*(.+)$", re.I), "TP. "),
    (re.compile(r"^(?:thị xã|thi xa|tx)\s*\.?\s*(.+)$", re.I), "TX. "),
    (re.compile(r"^(?:huyện|huyen)\s*\.?\s*(.+)$|^h\.\s*(.+)$|^h\s+(.+)$", re.I), "Huyện "),
    (re.compile(r"^(?:quận|quan)\s*\.?\s*(.+)$|^q\.\s*(.+)$|^q\s+(.+)$", re.I), "Quận "),
]


def district(value) -> str:
    """Chuẩn hóa tên quận/huyện: 'Q.Tân Bình' -> 'Quận Tân Bình', 'Q6' -> 'Quận 6', 'Thủ Đức' -> 'TP. Thủ Đức'."""
    s = _nfc(value)
    if not s:
        return UNKNOWN
    m = _RE_QUAN_SO.match(s) or re.fullmatch(r"0*(\d{1,2})", s)
    if m:
        return f"Quận {int(m.group(1))}"
    prefix, rest = "", s
    for rx, pre in _PREFIXES:
        m = rx.match(s)
        if m:
            prefix, rest = pre, next(g for g in m.groups() if g)
            break
    rest = _title(rest)
    if re.fullmatch(r"0*\d{1,2}", rest):
        return f"Quận {int(rest)}"
    k = plain(rest)
    if k in HUYEN:
        return f"Huyện {rest}"
    if k in THANH_PHO:
        return f"TP. {rest}"
    if k in QUAN_TEN:
        return f"Quận {rest}"
    return f"{prefix}{rest}" if prefix else rest


def ward(value) -> str:
    """'P.15' -> 'Phường 15', 'P. Tây Thạnh' -> 'Phường Tây Thạnh', 'X. Vĩnh Lộc A' -> 'Xã Vĩnh Lộc A'."""
    s = _nfc(value)
    if not s:
        return UNKNOWN
    for rx, pre in [(r"^(?:phường|phuong|p)\s*\.?\s*(.+)$", "Phường "),
                    (r"^(?:thị trấn|thi tran|tt)\s*\.?\s*(.+)$", "Thị trấn "),
                    (r"^(?:xã|xa|x)\s*\.\s*(.+)$|^(?:xã|xa)\s+(.+)$", "Xã ")]:
        m = re.match(rx, s, re.I)
        if m:
            rest = _title(next(g for g in m.groups() if g))
            return pre + (str(int(rest)) if rest.isdigit() else rest)
    if s.isdigit():
        return f"Phường {int(s)}"
    return _title(s)


def province(value) -> str:
    s = _nfc(value)
    if not s:
        return UNKNOWN
    k = plain(s).replace(".", " ")
    for name, label in TINH_LON.items():
        if name in k or name.replace(" ", "") in k.replace(" ", ""):
            return label
    s = re.sub(r"^(?:tỉnh|tinh)\s*\.?\s*", "", s, flags=re.I)
    return _title(s)


def _is_province(token: str) -> bool:
    k = plain(token)
    return k.startswith(("tinh ", "tinh.")) or any(n in k or n.replace(" ", "") in k.replace(" ", "")
                                                    for n in TINH_LON)


def parse_address(text) -> tuple[str, str, str]:
    """Tách (phường/xã, quận/huyện, tỉnh/TP) từ địa chỉ dạng '12 Lũy Bán Bích, P. Tân Thới Hòa, Q. Tân Phú, TP.HCM'."""
    parts = [p.strip() for p in re.split(r"[,;]", _nfc(text)) if p.strip()]
    w = q = t = ""
    used = set()
    for i in range(len(parts) - 1, -1, -1):
        part, k = parts[i], plain(parts[i])
        if not t and _is_province(part):
            t, _ = province(part), used.add(i)
        elif not q and (re.match(r"^(quan|q\.?\s*\d|q\.|q |huyen|h\.|thi xa|tx\.?|tp\.?|thanh pho)", k)
                        or k in HUYEN | QUAN_TEN | THANH_PHO):
            q, _ = district(part), used.add(i)
        elif not w and re.match(r"^(phuong|p\.?\s*\d|p\.|xa |x\.|thi tran|tt\.?)", k):
            w, _ = ward(part), used.add(i)
    # "..., P. Lái Thiêu, TP. Thuận An, Bình Dương": phần cuối chưa nhận diện được là tỉnh
    if not t and q and len(parts) >= 3 and len(parts) - 1 not in used and not re.search(r"\d", parts[-1]):
        t = province(parts[-1])
    return w, q, t


# ---- Khối / cấp ----
def grade(value) -> str:
    s = _nfc(value)
    if not s:
        return UNKNOWN
    m = re.search(r"(?<!\d)(1[0-2]|[1-9])(?!\d)", s)
    if m:
        return f"Khối {int(m.group(1))}"
    k = plain(s)
    if "mam non" in k or k in ("mn", "la", "choi", "mam"):
        return "Mầm non"
    return _title(s)


def level_of(khoi: str) -> str:
    m = re.search(r"\d+", khoi or "")
    if not m:
        return "Mầm non" if khoi == "Mầm non" else UNKNOWN
    n = int(m.group())
    return "Tiểu học" if n <= 5 else "THCS" if n <= 9 else "THPT"


def grade_order(label: str) -> tuple:
    m = re.search(r"\d+", label or "")
    return (0, int(m.group())) if m else (1 if label != UNKNOWN else 2, label)


# ---------------------------------------------------------------------------
# Chuẩn bị dữ liệu
# ---------------------------------------------------------------------------
def prepare(df: pd.DataFrame, cols: dict[str, str | None]) -> pd.DataFrame:
    """Bảng chuẩn: mỗi dòng một hồ sơ với các cột Nam, Tinh, QuanHuyen, PhuongXa, Khoi, Cap, GioiTinh, Nguon, TrangThai."""
    n = len(df)

    def col(field) -> pd.Series:
        name = cols.get(field)
        return df[name] if name and name in df.columns else pd.Series([None] * n, index=df.index, dtype=object)

    out = pd.DataFrame(index=df.index)
    nam = col("nam").map(school_year)
    if cols.get("ngay"):
        nam = nam.where(nam != "", year_from_date(col("ngay")))
    out["Nam"] = nam.replace("", UNKNOWN)

    raw_q, raw_w, raw_t = col("quan").map(_nfc), col("phuong").map(_nfc), col("tinh").map(_nfc)
    if cols.get("diachi"):
        parsed = col("diachi").map(parse_address)
        raw_w = raw_w.where(raw_w != "", parsed.map(lambda x: x[0]))
        raw_q = raw_q.where(raw_q != "", parsed.map(lambda x: x[1]))
        raw_t = raw_t.where(raw_t != "", parsed.map(lambda x: x[2]))
    # Tên trùng nhau được chuẩn hóa một lần (nhanh với vài chục nghìn dòng)
    out["QuanHuyen"] = raw_q.map({v: district(v) for v in raw_q.unique()})
    out["PhuongXa"] = raw_w.map({v: ward(v) for v in raw_w.unique()})
    out["Tinh"] = raw_t.map({v: province(v) for v in raw_t.unique()})
    out["Khoi"] = col("khoi").map(grade)
    cap = col("cap").map(_nfc)
    out["Cap"] = cap.where(cap != "", out["Khoi"].map(level_of)).replace("", UNKNOWN)
    for field, name in [("gioitinh", "GioiTinh"), ("nguon", "Nguon"), ("trangthai", "TrangThai")]:
        out[name] = col(field).map(_nfc).replace("", UNKNOWN)
    return out


def fill_province(*frames: pd.DataFrame) -> None:
    """Điền Tỉnh/TP còn trống theo quận/huyện (lấy tỉnh hay gặp nhất của quận đó ở cả hai list),
    không có thì lấy tỉnh phổ biến nhất – để một quận không bị tách thành 2 dòng."""
    both = pd.concat([f[["QuanHuyen", "Tinh"]] for f in frames])
    known = both[both["Tinh"] != UNKNOWN]
    if known.empty:
        return
    by_q = known.groupby("QuanHuyen")["Tinh"].agg(lambda s: s.value_counts().index[0])
    default = known["Tinh"].value_counts().index[0]
    for f in frames:
        miss = f["Tinh"] == UNKNOWN
        f.loc[miss, "Tinh"] = f.loc[miss, "QuanHuyen"].map(by_q).fillna(default)


def years_of(*frames: pd.DataFrame) -> list[str]:
    ys = set()
    for f in frames:
        ys |= set(f["Nam"])
    return sorted(ys, key=lambda y: (year_start(y), y))


def _pct(a, b):
    return round(a / b * 100, 1) if b else None


# ---------------------------------------------------------------------------
# Bảng thống kê
# ---------------------------------------------------------------------------
def by_year(data: pd.DataFrame, nh: pd.DataFrame, years: list[str]) -> pd.DataFrame:
    d, e = data["Nam"].value_counts(), nh["Nam"].value_counts()
    rows, prev = [], None
    for y in years:
        nd, ne = int(d.get(y, 0)), int(e.get(y, 0))
        rows.append({"Nam": y, "Data": nd, "NhapHoc": ne, "TyLe": _pct(ne, nd),
                     "TangGiam": (round((ne - prev) / prev * 100, 1) if prev else None)})
        prev = ne
    df = pd.DataFrame(rows)
    if not df.empty:
        df = pd.concat([df, pd.DataFrame([{"Nam": TOTAL, "Data": df["Data"].sum(), "NhapHoc": df["NhapHoc"].sum(),
                                           "TyLe": _pct(df["NhapHoc"].sum(), df["Data"].sum()), "TangGiam": None}])],
                       ignore_index=True)
    return df


def pivot(df: pd.DataFrame, dim: str, years: list[str], top: int | None = None) -> pd.DataFrame:
    """Dòng = ``dim``, cột = các năm học + Tổng; có dòng Tổng cộng."""
    if df.empty:
        return pd.DataFrame(columns=[DIMS.get(dim, dim), *years, "Tổng"])
    p = pd.crosstab(df[dim], df["Nam"]).reindex(columns=years, fill_value=0)
    p["Tổng"] = p.sum(axis=1)
    if dim == "Khoi":
        p = p.loc[sorted(p.index, key=grade_order)]
    else:
        p = p.sort_values("Tổng", ascending=False)
    if top and len(p) > top:
        rest = p.iloc[top:].sum()
        p = pd.concat([p.iloc[:top], pd.DataFrame([rest], index=[f"Khác ({len(p) - top})"])])
    p.loc[TOTAL] = p.sum()
    p.columns.name = None
    return p.reset_index(names=DIMS.get(dim, dim)).astype({c: int for c in [*years, "Tổng"]})


def compare(data: pd.DataFrame, nh: pd.DataFrame, dim: str, years: list[str], group: str | None = None) -> pd.DataFrame:
    """Mẫu chi tiết: dòng = ``dim`` (vd Quận/Huyện), mỗi năm 3 cột Data · Nhập học · Tỷ lệ %, cuối cùng là Tổng.
    ``group`` (vd Tỉnh/TP) thêm cột nhóm phía trước."""
    keys = [group, dim] if group else [dim]
    d = data.groupby(keys + ["Nam"]).size().unstack("Nam").reindex(columns=years).fillna(0)
    e = nh.groupby(keys + ["Nam"]).size().unstack("Nam").reindex(columns=years).fillna(0)
    idx = d.index.union(e.index)
    d, e = d.reindex(idx, fill_value=0), e.reindex(idx, fill_value=0)
    out = pd.DataFrame(index=idx)
    for y in [*years, "Tổng"]:
        dv = d.sum(axis=1) if y == "Tổng" else d[y]
        ev = e.sum(axis=1) if y == "Tổng" else e[y]
        out[f"{y} · {DATA}"] = dv.astype(int)
        out[f"{y} · {NH}"] = ev.astype(int)
        out[f"{y} · Tỷ lệ %"] = [(_pct(a, b)) for a, b in zip(ev, dv)]
    out = out.sort_values([f"Tổng · {NH}", f"Tổng · {DATA}"], ascending=False)
    if dim == "Khoi":
        out = out.loc[sorted(out.index, key=lambda k: grade_order(k[-1] if isinstance(k, tuple) else k))]
    out = out.reset_index()
    out.columns = [DIMS.get(c, c) for c in out.columns]
    if not out.empty:
        total = {c: out[c].sum() for c in out.columns if c.endswith((DATA, NH))}
        for y in [*years, "Tổng"]:
            total[f"{y} · Tỷ lệ %"] = _pct(total[f"{y} · {NH}"], total[f"{y} · {DATA}"])
        total[DIMS[dim]] = TOTAL
        if group:
            total[DIMS[group]] = ""
        out = pd.concat([out, pd.DataFrame([total])], ignore_index=True)
    return out


# ---------------------------------------------------------------------------
# Excel: mẫu chi tiết theo quận/huyện
# ---------------------------------------------------------------------------
def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Sheet"
    cand, i = base, 2
    while cand.lower() in used:
        cand = f"{base[:28]} {i}"
        i += 1
    used.add(cand.lower())
    return cand


def _write_compare(ws, top: int, table: pd.DataFrame, years: list[str], label_cols: list[str]) -> int:
    """Ghi bảng ``compare`` có tiêu đề 2 tầng (Năm học / Data · Nhập học · Tỷ lệ) từ dòng ``top``; trả về dòng kế tiếp."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DCE6F5")
    total_fill = PatternFill("solid", fgColor="F2F2F2")
    bold = Font(name="Times New Roman", size=11, bold=True)
    normal = Font(name="Times New Roman", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    heads = ["STT", *label_cols]
    for i, h in enumerate(heads, start=1):
        ws.merge_cells(start_row=top, start_column=i, end_row=top + 1, end_column=i)
        c = ws.cell(top, i, h)
    col = len(heads) + 1
    for y in [*years, "Tổng"]:
        ws.merge_cells(start_row=top, start_column=col, end_row=top, end_column=col + 2)
        ws.cell(top, col, "Tổng các năm" if y == "Tổng" else f"Năm học {y}")
        for j, sub in enumerate(["Data TS", "Nhập học", "Tỷ lệ (%)"]):
            ws.cell(top + 1, col + j, sub)
        col += 3
    last_col = col - 1
    for r in (top, top + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font, cell.alignment, cell.border, cell.fill = bold, center, border, head_fill

    row = top + 2
    stt = 0
    for _, rec in table.iterrows():
        is_total = rec[label_cols[-1]] == TOTAL
        if not is_total:
            stt += 1
        ws.cell(row, 1, "" if is_total else stt)
        for i, lc in enumerate(label_cols, start=2):
            ws.cell(row, i, rec[lc])
        col = len(heads) + 1
        for y in [*years, "Tổng"]:
            for j, suffix in enumerate([DATA, NH, "Tỷ lệ %"]):
                v = rec[f"{y} · {suffix}"]
                ws.cell(row, col + j, None if pd.isna(v) else float(v) if j == 2 else int(v))
                ws.cell(row, col + j).number_format = "0.0" if j == 2 else "#,##0"
            col += 3
        for c in range(1, last_col + 1):
            cell = ws.cell(row, c)
            cell.border = border
            cell.font = bold if is_total else normal
            if is_total:
                cell.fill = total_fill
            if c == 1 or c > len(heads):
                cell.alignment = Alignment(horizontal="center" if c == 1 else "right", vertical="center")
        row += 1
    return row


def _sheet_header(ws, school: str, title: str, subtitle: str, width: int) -> int:
    from openpyxl.styles import Alignment, Font

    ws.cell(1, 1, school).font = Font(name="Times New Roman", size=12, bold=True)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(width, 4))
    c = ws.cell(3, 1, title)
    c.font = Font(name="Times New Roman", size=15, bold=True)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max(width, 4))
    c = ws.cell(4, 1, subtitle)
    c.font = Font(name="Times New Roman", size=11, italic=True)
    c.alignment = Alignment(horizontal="center")
    return 6


def _signature(ws, row: int, width: int) -> None:
    from openpyxl.styles import Alignment, Font

    today = date.today()
    left, right = 2, max(width - 3, 5)
    ws.cell(row + 1, right, f"TP. Hồ Chí Minh, ngày {today.day} tháng {today.month} năm {today.year}").font = \
        Font(name="Times New Roman", size=11, italic=True)
    for col, text in [(left, "NGƯỜI LẬP BIỂU"), (right, "BAN GIÁM HIỆU")]:
        c = ws.cell(row + 2, col, text)
        c.font = Font(name="Times New Roman", size=11, bold=True)
        c.alignment = Alignment(horizontal="center")


def _widths(ws, label_cols: int, n_years: int) -> None:
    from openpyxl.utils import get_column_letter

    ws.column_dimensions["A"].width = 6
    for i in range(label_cols):
        ws.column_dimensions[get_column_letter(2 + i)].width = 26
    for i in range((n_years + 1) * 3):
        ws.column_dimensions[get_column_letter(2 + label_cols + i)].width = 10


def district_xlsx(data: pd.DataFrame, nh: pd.DataFrame, years: list[str], school: str, note: str = "",
                  detail: bool = True) -> bytes:
    """Mẫu chi tiết theo quận/huyện: sheet tổng hợp + mỗi quận/huyện một sheet (theo phường/xã và theo khối)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    used: set[str] = set()
    ws.title = _sheet_name("Tổng hợp quận-huyện", used)
    multi_tinh = len((set(data["Tinh"]) | set(nh["Tinh"])) - {UNKNOWN}) > 1
    group = "Tinh" if multi_tinh else None
    table = compare(data, nh, "QuanHuyen", years, group=group)
    labels = [DIMS["Tinh"], DIMS["QuanHuyen"]] if group else [DIMS["QuanHuyen"]]
    width = 1 + len(labels) + 3 * (len(years) + 1)
    sub = f"Năm học {years[0]} – {years[-1]}" if years else ""
    top = _sheet_header(ws, school, "BÁO CÁO TUYỂN SINH THEO QUẬN/HUYỆN QUA CÁC NĂM",
                        sub + (f" · {note}" if note else ""), width)
    end = _write_compare(ws, top, table, years, labels)
    _signature(ws, end + 1, width)
    _widths(ws, len(labels), len(years))
    ws.freeze_panes = ws.cell(top + 2, len(labels) + 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    if detail:
        districts = [q for q in table[DIMS["QuanHuyen"]] if q not in (TOTAL, "")]
        for q in dict.fromkeys(districts):
            d, e = data[data["QuanHuyen"] == q], nh[nh["QuanHuyen"] == q]
            ws = wb.create_sheet(_sheet_name(q, used))
            width = 2 + 3 * (len(years) + 1)
            top = _sheet_header(ws, school, f"CHI TIẾT TUYỂN SINH – {q.upper()}", sub, width)
            ws.cell(top - 1, 1, "1. Theo phường/xã").font = _bold_font()
            end = _write_compare(ws, top, compare(d, e, "PhuongXa", years), years, [DIMS["PhuongXa"]])
            ws.cell(end + 1, 1, "2. Theo khối").font = _bold_font()
            end = _write_compare(ws, end + 2, compare(d, e, "Khoi", years), years, [DIMS["Khoi"]])
            _signature(ws, end + 1, width)
            _widths(ws, 1, len(years))
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bold_font():
    from openpyxl.styles import Font

    return Font(name="Times New Roman", size=12, bold=True)
