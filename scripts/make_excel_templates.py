"""Tạo file Excel mẫu cho từng SharePoint List (dùng "Microsoft Lists → Danh sách mới → Từ Excel").

    python scripts/make_excel_templates.py   # ghi vào thư mục templates/

Tên cột trong file = tên hiển thị mà ứng dụng dùng để dò cột, nên list tạo từ file sẽ kết nối được ngay.
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from qlts import schema  # noqa: E402

OUT = ROOT / "templates"

TYPE_NAMES = {
    "text": "Một dòng văn bản (Single line of text)",
    "note": "Nhiều dòng văn bản (Multiple lines of text)",
    "number": "Số (Number)",
    "date": "Ngày (Date and time – chỉ ngày)",
    # Tình trạng / Vai trò để dạng văn bản cho linh hoạt (dữ liệu cũ có nhiều giá trị khác nhau)
    "choice": "Một dòng văn bản (Single line of text)",
}

SAMPLES = {
    schema.THIET_BI: {
        "STT": 1, "MaTaiSan": "111028", "MaChiTiet": "111028-00001", "ChiTiet": "Máy chiếu",
        "DacDiem": "Epson, 1024x768, EB-530", "TenPhongBan": "VFQF610018L", "DVT": "Cái", "SL": 1,
        "NoiSuDung": "Trệt_PH001", "NguoiSuDung": "giaovien@truong.edu.vn", "NgayMua": date(2023, 8, 22),
        "QuanLyThietBi": "Lê Thiên An", "TinhTrang": "Bình thường", "QuanLyPhong": "giaovien@truong.edu.vn",
        "PhanQuyenTB": "admin", "MaSAP": "T000816001", "ThoiHanBaoHanh": "24 tháng", "GhiChu": "",
        "TenThietBi": "Máy chiếu", "GiaTri": 12000000, "NgayHoaDon": date(2023, 8, 22),
        "NhomThietBi": "Nhóm Công Nghệ Thông Tin", "Mail": "",
    },
    schema.PHONG: {
        "Title": "Trệt_PH001", "TenPhong": "Phòng học 001", "KhuVuc": "Tầng trệt",
        "NguoiQuanLy": "giaovien@truong.edu.vn", "TenNguoiQuanLy": "Nguyễn Văn A",
    },
    schema.DIEU_CHUYEN: {
        "Title": "111028-00001", "TenThietBi": "Máy chiếu", "TuPhong": "Kho_TB", "DenPhong": "Trệt_PH001",
        "SoLuong": 1, "NgayDieuChuyen": date(2026, 9, 1), "NguoiThucHien": "qlts@truong.edu.vn",
        "LyDo": "Cấp cho phòng học",
    },
    schema.KIEM_KE: {
        "MaTaiSan": "111042", "STT": 1, "TenTaiSan": "Tivi", "DacDiem": "Sharp, 58 inch, UE630X", "DVT": "Cái",
        "SoLuong": 1, "SoLuongKiemKe": 1, "GhiChu": "", "QuanLyThietBi": "Lê Thiên An", "NoiSuDung": "Bảo vệ",
        "QuanLyPhong": "Nguyễn Trọng Hòa", "TrangThai": "Bình thường", "NgayKiemKe": date(2025, 11, 5),
        "TrangThaiKiemKe": "Đã kiểm kê", "NgayXacNhan": None, "TrangThaiXacNhan": "",
        "DotKiemKe": "Đợt 2 năm 2024 - 2025", "NguoiKiemKe": "kiemke@truong.edu.vn", "NguoiXacNhan": "",
    },
    schema.PHAN_QUYEN: {
        "Title": "email_cua_ban@truong.edu.vn", "HoTen": "Lê Thiên An", "VaiTro": schema.ROLE_ADMIN,
        "ChucDanh": "Chuyên viên Quản lý hệ thống",
    },
}

NOTES = {
    "MaChiTiet": "Mã tài sản + '-' + 5 chữ số, app tự sinh khi nhập mới",
    "NoiSuDung": "Mã phòng (trùng cột Mã phòng của list Phong); 'Thanh lý' = đã đưa đi thanh lý",
    "QuanLyPhong": "Email người quản lý phòng – người này thấy phòng ở mục Người dùng",
    "TinhTrang": "Gợi ý: " + ", ".join(schema.TINH_TRANG),
    "VaiTro": "Một trong: " + ", ".join(schema.ROLES) + " (mỗi vai trò một dòng)",
    "SoLuong": "Số lượng sổ sách = tổng thiết bị cùng Mã tài sản + Đặc điểm + Nơi sử dụng",
    "TrangThaiKiemKe": "Đang kiểm kê / Đã kiểm kê (app ghi khi bấm Xác nhận kiểm kê)",
    "TrangThaiXacNhan": "Đã xác nhận / Không đồng ý (quản lý phòng xác nhận)",
    "NguoiKiemKe": "Không bắt buộc – có cột thì app ghi email người kiểm kê",
    "NguoiXacNhan": "Không bắt buộc – có cột thì app ghi email người xác nhận",
    "Title": "",
}

FILES = {
    schema.THIET_BI: ("1_ThietBi.xlsx", "Data_Thietbichitiet (hoặc tên bạn chọn)"),
    schema.PHONG: ("2_Phong.xlsx", "Phong"),
    schema.DIEU_CHUYEN: ("3_DieuChuyen.xlsx", "DieuChuyen"),
    schema.KIEM_KE: ("4_KiemKe.xlsx", "Data_Thietbi"),
    schema.PHAN_QUYEN: ("5_PhanQuyen.xlsx", "PhanQuyen"),
}


def build(list_key: str, path: Path, list_name: str) -> None:
    cols = schema.LISTS[list_key]["columns"]
    wb = Workbook()
    ws = wb.active
    ws.title = "DuLieu"
    ws.append([spec["label"] for spec in cols.values()])
    sample = SAMPLES[list_key]
    ws.append([sample.get(key, "") for key in cols])
    for i, (key, spec) in enumerate(cols.items(), start=1):
        cell = ws.cell(row=2, column=i)
        if spec["type"] == "date":
            cell.number_format = "DD/MM/YYYY"
        elif spec["type"] == "number":
            cell.number_format = "#,##0"
        ws.column_dimensions[cell.column_letter].width = max(14, len(spec["label"]) + 4)
    table = Table(displayName=f"tbl_{list_key}", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    guide = wb.create_sheet("HuongDan")
    guide.append([f"List SharePoint: {list_name}"])
    guide["A1"].font = Font(bold=True, size=13)
    guide.append(["Tạo list: Microsoft Lists → + Danh sách mới → Từ Excel → chọn file này → chọn bảng "
                  f"tbl_{list_key} → kiểm tra kiểu cột theo bảng dưới → Tạo. Sau đó XÓA dòng mẫu."])
    guide.append(["Không đổi tên cột: app dò cột theo đúng tên này."])
    guide.append([])
    guide.append(["Tên cột", "Kiểu cột khi tạo list", "Ghi chú"])
    for c in guide[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E8C")
    for key, spec in cols.items():
        guide.append([spec["label"], TYPE_NAMES[spec["type"]], NOTES.get(key, "")])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 44
    guide.column_dimensions["C"].width = 70
    for row in guide.iter_rows(min_row=2, max_row=3):
        row[0].alignment = Alignment(wrap_text=False)
    wb.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for key, (filename, list_name) in FILES.items():
        build(key, OUT / filename, list_name)
        print("Đã tạo", OUT / filename)


if __name__ == "__main__":
    main()
